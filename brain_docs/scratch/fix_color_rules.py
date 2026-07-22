import sys
import re
import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding='utf-8')

f1 = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720.xlsx'
f2 = r'C:\Users\Administrator\Desktop\배현명만세.xlsx'
f2_alt = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명만세.xlsx'

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid") # 신규 채움 (노란색)
orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid") # 수정 입력 (주황색)
no_fill = PatternFill(fill_type=None) # 색상 없음 (삭제건 및 의미 동일건)

def normalize_val(val, col_name=""):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str == "" or val_str == "-":
        return ""
    col_name = str(col_name)
    if any(k in col_name for k in ['전화', '팩스', '사업자', '우편', '번호', '일자', '일']):
        clean_num = re.sub(r'[^0-9]', '', val_str)
        if clean_num:
            return clean_num
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

print("1. 엑셀 파일 로딩 중...")
wb1 = openpyxl.load_workbook(f1, data_only=True, read_only=True)
wb2 = openpyxl.load_workbook(f2)

s1 = wb1['기업 정보']
s2 = wb2['기업 정보']

max_c = 60
s1_rows = list(s1.iter_rows(max_col=max_c, values_only=True))
headers1 = [str(h).strip() if h is not None else "" for h in s1_rows[0]]
headers2 = [str(cell.value).strip() if cell.value is not None else "" for cell in list(s2.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]]

col_map2 = {name: idx+1 for idx, name in enumerate(headers2)}

biz_c = col_map2.get('사업자번호')
comp_c = col_map2.get('회사명')
author_c = col_map2.get('작성자')

company_idx1 = headers1.index('회사명') if '회사명' in headers1 else -1
biz_idx1 = headers1.index('사업자번호') if '사업자번호' in headers1 else -1

col_mapping = {}
for idx2, h2 in enumerate(headers2):
    if h2 in headers1:
        col_mapping[idx2 + 1] = headers1.index(h2)

f1_key_map = {}
for r_idx, r_val in enumerate(s1_rows[1:], start=2):
    biz = str(r_val[biz_idx1]).strip() if biz_idx1 != -1 and biz_idx1 < len(r_val) and r_val[biz_idx1] else ""
    comp = str(r_val[company_idx1]).strip() if company_idx1 != -1 and company_idx1 < len(r_val) and r_val[company_idx1] else ""
    clean_b = re.sub(r'[^0-9]', '', biz)
    key = (clean_b, comp)
    if key != ("", "") and key not in f1_key_map:
        f1_key_map[key] = r_val

print("2. 색상 재정의 규칙 적용 (삭제 셀 색상 안 칠함)...")

yellow_count = 0
orange_count = 0
deleted_count = 0
unchanged_count = 0

for r_idx2, row_cells in enumerate(s2.iter_rows(min_row=2, max_col=max_c), start=2):
    val_author = str(row_cells[author_c-1].value).strip() if author_c and row_cells[author_c-1].value is not None else ""
    
    if '배현명(7.20~)' not in val_author:
        continue
        
    biz2 = ""
    comp2 = ""
    for c_idx2, cell2 in enumerate(row_cells, start=1):
        h2_name = headers2[c_idx2 - 1] if c_idx2 - 1 < len(headers2) else ""
        if h2_name == '사업자번호':
            biz2 = normalize_val(cell2.value, '사업자번호')
        elif h2_name == '회사명':
            comp2 = normalize_val(cell2.value, '회사명')
            
    key2 = (biz2, comp2)
    
    target_r1_val = None
    if key2 != ("", "") and key2 in f1_key_map:
        target_r1_val = f1_key_map[key2]
    elif r_idx2 - 1 < len(s1_rows):
        target_r1_val = s1_rows[r_idx2 - 1]
        
    if target_r1_val is None:
        for c_idx2, cell2 in enumerate(row_cells, start=1):
            col_name = headers2[c_idx2 - 1] if c_idx2 - 1 < len(headers2) else ""
            n_val2 = normalize_val(cell2.value, col_name)
            if n_val2 != "":
                cell2.fill = yellow_fill
                yellow_count += 1
            else:
                cell2.fill = no_fill
        continue
        
    for c_idx2, cell2 in enumerate(row_cells, start=1):
        if c_idx2 not in col_mapping:
            continue
        c_idx1 = col_mapping[c_idx2]
        col_name = headers2[c_idx2 - 1] if c_idx2 - 1 < len(headers2) else ""
        
        orig_raw_val = target_r1_val[c_idx1] if c_idx1 < len(target_r1_val) else None
        new_raw_val = cell2.value
        
        norm_orig = normalize_val(orig_raw_val, col_name)
        norm_new = normalize_val(new_raw_val, col_name)
        
        # 1. 신규 입력 -> 🟡 노란색
        if norm_orig == "" and norm_new != "":
            cell2.fill = yellow_fill
            yellow_count += 1
            
        # 2. 수정 입력 (원래 다른 값이 있었는데 새로운 값으로 변경됨) -> 🟠 주황색
        elif norm_orig != "" and norm_new != "":
            if norm_orig != norm_new:
                cell2.fill = orange_fill
                orange_count += 1
            else:
                cell2.fill = no_fill
                unchanged_count += 1
                
        # 3. 삭제된 셀 (원래 값이 있었는데 지워져서 비워짐) -> ⚪ 색상 안 칠함 (no_fill)
        elif norm_orig != "" and norm_new == "":
            cell2.fill = no_fill
            deleted_count += 1
            
        else:
            cell2.fill = no_fill

print(f"\n==========================================")
print(f"🎨 수정된 색상 규칙 적용 결과")
print(f"==========================================")
print(f" - 🟡 신규 입력 (노란색 셀): {yellow_count}개")
print(f" - 🟠 수정 입력 (주황색 셀): {orange_count}개")
print(f" - ⚪ 삭제된 셀 (색상 안 칠함): {deleted_count}개")
print(f" - ⚪ 의미상 동일/미변경 (색상 안 칠함): {unchanged_count}개")
print(f"==========================================")

print("3. 결과 파일 저장 중...")
wb2.save(f2)
wb2.save(f2_alt)
print(f"✅ 저장 완료:\n 1. {f2}\n 2. {f2_alt}")

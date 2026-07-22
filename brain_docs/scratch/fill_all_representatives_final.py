import sys
import re
import urllib.request
import urllib.parse
import concurrent.futures
import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding='utf-8')

f_in = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_검증완료.xlsx'
f_out = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_검증완료.xlsx'

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

biz_list = [
    ('6068540092', '(주)거보녹산지점'), ('8128501734', '(주)건웅기계송정동지점'), ('7498500515', '(주)경동특수주강화전지점'),
    ('1018549002', '(주)까사세라믹홈디포'), ('6068526839', '(주)대아선재녹산공장'), ('4738500113', '(주)대천기장지점'),
    ('6218523656', '(주)동양테크정관'), ('2118575637', '(주)베르디-베르디사우나'), ('6068545404', '(주)비앤에스푸드금문산업'),
    ('3518500263', '(주)삼락열처리미음지점'), ('6038507435', '(주)성우하이텍지사공장'), ('6648500271', '(주)성원화학상사녹산지점'),
    ('8478500390', '(주)세진ELT'), ('8488500689', '(주)손안기전'), ('6068544477', '(주)아산부산지점'),
    ('6068538335', '(주)염성테크부산지점'), ('8488501183', '(주)웨스코일렉트로드'), ('7208500339', '(주)유성신소재화전지점'),
    ('3608500350', '(주)준코-경성대점'), ('6068539616', '(주)천우테크녹산공장'), ('8688501741', '(주)쿤스트원부산센텀지점'),
    ('5918501856', '(주)클루커스부산캠퍼스'), ('6178533868', '(주)푸드엔경성대점'), ('5958501988', '(주)푸드엔남천지점'),
    ('3818501657', '(주)푸드엔엄궁지점'), ('6068547760', '(주)한일경판2공장'), ('6218536092', '(주)한텍테크놀로지장안지점'),
    ('6068526047', '건우메탈(주)'), ('7988501893', '국제커리어센터부산남구'), ('2638500855', '남양엔지니어링주식회사녹산지점'),
    ('6068541198', '대신노무법인(송정)'), ('3053451388', '더퍼스트피부과의원'), ('6068548338', '동서열처리(주)미음지점'),
    ('6068515399', '동아흥업(주)부산공장'), ('6743000109', '동현테크'), ('1518500722', '두끼CGV대연점'),
    ('5585201074', '미주물류'), ('6068525150', '비엔철강(주)코일센타'), ('6068542130', '삼영엠티(주)지사지점'),
    ('6218524766', '삼원산업(주)부산공장'), ('6068546043', '세무법인부경사상지점'), ('6978501912', '세무법인부성남부산지점'),
    ('3668501226', '세무법인서초부산센터'), ('6068543783', '세무법인시우(북부산지점)'), ('6068539111', '세무법인중추북부산지점'),
    ('6178532137', '세무법인태율'), ('2788600868', '스카이오토(주)'), ('7288501763', '신한테크주식회사지점'),
    ('4868500991', '에스엠정밀기술'), ('6178530819', '에스피앤지(주)기업부설연구소'), ('7888501464', '정아정밀(주)강서공장'),
    ('7348500677', '존슨콘트롤즈인터내셔널코리아㈜조선마린지점'), ('1778502232', '주식회사등푸른식품선진화단지점'),
    ('6178620720', '주식회사미옹'), ('5788602432', '주식회사에이스로직스'), ('6348500232', '주식회사코리아테크화전지점'),
    ('6068525922', '주식회사풍진금속'), ('3354801343', '중부기업'), ('3478102890', '태상산업주식회사'),
    ('7288502272', '파이널나인부산경성대점'), ('7268501224', '한국코벨코용접(주)부산지점')
]

def fetch_rep_name(biz, comp):
    url = f'https://bizno.net/article/{biz}'
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
    for attempt in range(2):
        try:
            with urllib.request.urlopen(req, timeout=5) as resp:
                html = resp.read().decode('utf-8', errors='ignore')
                items = re.findall(r'<th>(.*?)</th>\s*<td>(.*?)</td>', html, re.DOTALL)
                d = {}
                for th, td in items:
                    th_clean = re.sub(r'<[^>]+>', '', th).strip()
                    td_clean = re.sub(r'<[^>]+>', '', td).strip()
                    d[th_clean] = td_clean
                rep = d.get('대표자명') or d.get('대표자')
                if rep:
                    return biz, comp, rep
        except Exception:
            pass
    return biz, comp, None

print("1. 비즈노 대표자 정보 수집 중 (멀티 스레드)...")
rep_results = {}
with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
    futures = [executor.submit(fetch_rep_name, biz, comp) for biz, comp in biz_list]
    for future in concurrent.futures.as_completed(futures):
        biz, comp, rep = future.result()
        if rep:
            rep_results[biz] = rep
            print(f"  [성공] {comp} ({biz}) -> 대표자: {rep}")
        else:
            print(f"  [미발견/지점] {comp} ({biz})")

print(f"\n총 {len(rep_results)}개 기업의 대표자 성함 수집 완료.")

print("2. 엑셀 파일 반영 중...")
wb = openpyxl.load_workbook(f_in)
s = wb['기업 정보']

max_c = 60
headers = [str(cell.value).strip() if cell.value is not None else "" for cell in list(s.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]]
col_map = {name: idx+1 for idx, name in enumerate(headers)}

biz_c = col_map.get('사업자번호')
rep_c = col_map.get('대표자')
author_c = col_map.get('작성자')

filled_count = 0
for r_idx, row in enumerate(s.iter_rows(min_row=2, max_col=max_c), start=2):
    author = str(row[author_c-1].value) if author_c and row[author_c-1].value else ''
    if '배현명(7.20~)' in author:
        biz_val = str(row[biz_c-1].value).strip() if biz_c and row[biz_c-1].value else ''
        biz_clean = re.sub(r'[^0-9]', '', biz_val)
        
        if biz_clean in rep_results:
            cell_rep = row[rep_c-1]
            rep_name = rep_results[biz_clean]
            cell_rep.value = rep_name
            cell_rep.fill = yellow_fill
            filled_count += 1

print(f"3. 엑셀에 반영 완료: 총 {filled_count}개 대표자 노란색 셀 채움.")

try:
    wb.save(f_out)
    print(f"✅ 저장 완료: {f_out}")
except Exception as e:
    alt_out = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_검증완료_최종.xlsx'
    wb.save(alt_out)
    print(f"✅ 열려있는 파일 권한으로 인해 새 파일로 저장 완료: {alt_out}")

import sys
import re
import urllib.request
import urllib.parse
import json
import concurrent.futures
import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding='utf-8')

f_out = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_검증완료.xlsx'

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

def fetch_rep_from_bizno(biz_no, comp_name):
    clean_biz = re.sub(r'[^0-9]', '', str(biz_no)) if biz_no else ''
    rep_name = None
    phone_num = None
    fax_num = None
    
    # 1. 비즈노 직접 조회
    if clean_biz:
        url = f"https://bizno.net/article/{clean_biz}"
        try:
            req = urllib.request.Request(url, headers={
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
            })
            with urllib.request.urlopen(req, timeout=4) as resp:
                html = resp.read().decode('utf-8', errors='ignore')
                
                # 대표자 추출
                m = re.search(r'대표자[^\w\n]*</th>[^\w\n]*<td>\s*([^<\s]+(?:\s+[^<\s]+)?)\s*</td>', html)
                if m:
                    rep_name = m.group(1).strip()
                    
                # 전화번호 추출
                m_p = re.search(r'전화번호[^\w\n]*</th>[^\w\n]*<td>\s*([^<\n]+)\s*</td>', html)
                if m_p:
                    phone_num = m_p.group(1).strip()
                    
                # 팩스번호 추출
                m_f = re.search(r'팩스번호[^\w\n]*</th>[^\w\n]*<td>\s*([^<\n]+)\s*</td>', html)
                if m_f:
                    fax_num = m_f.group(1).strip()
        except Exception:
            pass
            
    # 2. 웹 검색 백업 (비즈노에서 안 나올 경우)
    if not rep_name and comp_name:
        try:
            query = f"{clean_biz} {comp_name} 대표자"
            search_url = f"https://html.duckduckgo.com/html/?q={urllib.parse.quote(query)}"
            req = urllib.request.Request(search_url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
            with urllib.request.urlopen(req, timeout=4) as resp:
                html = resp.read().decode('utf-8', errors='ignore')
                # 패턴 1: 대표자 : xxx 또는 대표 : xxx
                m = re.search(r'대표자?\s*[:\s]\s*([가-힣]{2,4})', html)
                if m:
                    candidate = m.group(1).strip()
                    if candidate not in ['주식회사', '부산', '강서구', '대표', '대표자', '기업']:
                        rep_name = candidate
        except Exception:
            pass
            
    return rep_name, phone_num, fax_num

print("1. 엑셀 파일 로딩 중...")
wb = openpyxl.load_workbook(f_out)
s = wb['기업 정보']

max_c = 60
headers = [str(cell.value).strip() if cell.value is not None else "" for cell in list(s.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]]

col_map = {name: idx+1 for idx, name in enumerate(headers)}

biz_c = col_map.get('사업자번호')
comp_c = col_map.get('회사명')
rep_c = col_map.get('대표자')
phone_c = col_map.get('전화번호')
fax_c = col_map.get('팩스번호')
author_c = col_map.get('작성자')

# 대표자가 비어있는 행 수집
target_rows = []
for r_idx, row in enumerate(s.iter_rows(min_row=2, max_col=max_c), start=2):
    author = str(row[author_c-1].value) if author_c and row[author_c-1].value else ''
    if '배현명(7.20~)' in author:
        rep_val = row[rep_c-1].value if rep_c else ''
        if not rep_val or str(rep_val).strip() in ['-', '']:
            biz = row[biz_c-1].value if biz_c else ''
            comp = row[comp_c-1].value if comp_c else ''
            target_rows.append((r_idx, biz, comp, row))

print(f"2. 대표자 공란 기업 총 {len(target_rows)}개 검색 및 보완 시작...")

def process_item(item):
    r_idx, biz, comp, row = item
    rep, phone, fax = fetch_rep_from_bizno(biz, comp)
    return r_idx, rep, phone, fax, row

rep_filled = 0
phone_filled = 0
fax_filled = 0

with concurrent.futures.ThreadPoolExecutor(max_workers=15) as executor:
    futures = [executor.submit(process_item, item) for item in target_rows]
    for future in concurrent.futures.as_completed(futures):
        r_idx, rep, phone, fax, row = future.result()
        
        # 대표자 채우기
        if rep and rep_c:
            cell_rep = row[rep_c-1]
            if not cell_rep.value or str(cell_rep.value).strip() in ['-', '']:
                cell_rep.value = rep
                cell_rep.fill = yellow_fill
                rep_filled += 1
                print(f"Row {r_idx} [{row[comp_c-1].value}] -> 대표자 채움: {rep}")
                
        # 전화번호 채우기
        if phone and phone_c:
            cell_phone = row[phone_c-1]
            if not cell_phone.value or str(cell_phone.value).strip() in ['-', '']:
                cell_phone.value = phone
                cell_phone.fill = yellow_fill
                phone_filled += 1
                
        # 팩스번호 채우기
        if fax and fax_c:
            cell_fax = row[fax_c-1]
            if not cell_fax.value or str(cell_fax.value).strip() in ['-', '']:
                cell_fax.value = fax
                cell_fax.fill = yellow_fill
                fax_filled += 1

print(f"\n3. 보완 완료 통계:")
print(f" - 대표자명 새로 채움: {rep_filled}개")
print(f" - 전화번호 새로 채움: {phone_filled}개")
print(f" - 팩스번호 새로 채움: {fax_filled}개")

print("4. 결과 저장 중...")
wb.save(f_out)
print(f"✅ 저장 완료: {f_out}")

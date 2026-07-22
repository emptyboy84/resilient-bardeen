import sys
import re
import os
import openpyxl
from openpyxl.styles import PatternFill, Font

sys.stdout.reconfigure(encoding='utf-8')

# 검증 및 유저 확인 완료된 실존 대표자명 매핑
real_rep_map = {
    '6068540092': '박임수',   # (주)거보녹산지점
    '8128501734': '박은영',   # (주)건웅기계송정동지점
    '7498500515': '오기호',   # (주)경동특수주강화전지점
    '1018549002': '조성원',   # (주)까사세라믹홈디포
    '6068526839': '진영준',   # (주)대아선재녹산공장
    '4738500113': '김병열',   # (주)대천기장지점
    '6218523656': '김정규',   # (주)동양테크정관
    '3518500263': '신호강',   # (주)삼락열처리미음지점
    '6038507435': '이명근',   # (주)성우하이텍지사공장
    '8488500689': '손기준',   # (주)손안기전
    '6068544477': '김예경',   # (주)아산부산지점
    '6178533868': '김광원',   # (주)푸드엔경성대점
    '5958501988': '김광원',   # (주)푸드엔남천지점
    '3818501657': '김광원',   # (주)푸드엔엄궁지점
    '6068541198': '신명석',   # 대신노무법인(송정)
    '3053451388': '공숙현',   # 더퍼스트피부과의원
    '6068548338': '신호강',   # 동서열처리(주)미음지점
    '6743000109': '김동옥',   # 동현테크
    '5585201074': '정상섭',   # 미주물류
    '3354801343': '서진수',   # 중부기업
    '2788600868': '오린태',   # 스카이오토(주)
}

possible_files = [
    r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\실사_DB_취합_전달용_배현명_업데이트.xlsx',
    r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx'
]

f_src = next((p for p in possible_files if os.path.exists(p)), None)
print(f"1. 로딩 대상 소스 파일: {f_src}")

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
no_fill = PatternFill(fill_type=None)
header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

def normalize_val(val, col_name=""):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str == "" or val_str == "-":
        return ""
    col_name = str(col_name)
    if any(k in col_name for k in ['전화', '팩스', '사업자', '우편', '번호', '일자', '일']):
        clean_num = re.sub(r'[^0-9]', '', val_str)
        if clean_num:
            return clean_num
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

wb = openpyxl.load_workbook(f_src)
s = wb['기업 정보'] if '기업 정보' in wb.sheetnames else wb.active

max_c = 60
headers = [str(cell.value).strip() if cell.value is not None else "" for cell in list(s.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]]

col_map = {name: idx+1 for idx, name in enumerate(headers)}

worker_c = col_map.get('작업자') or col_map.get('작성자') or 1
rep_c = col_map.get('대표자')
addr_c = col_map.get('주소')
phone_c = col_map.get('전화번호')
fax_c = col_map.get('팩스번호')
hp_c = col_map.get('홈페이지')
email_c = col_map.get('이메일')
emp_c = col_map.get('종업원수')
biz_c = col_map.get('사업자번호')
comp_c = col_map.get('회사명')

start_src_col = 58
src_headers = [
    '대표자_출처_FullURL',
    '주소_출처_FullURL',
    '전화번호_출처_FullURL',
    '팩스번호_출처_FullURL',
    '홈페이지_출처_FullURL',
    '이메일_출처_FullURL',
    '종업원수_출처_FullURL'
]

for i, h_title in enumerate(src_headers):
    c_idx = start_src_col + i
    cell = s.cell(row=1, column=c_idx)
    cell.value = h_title
    cell.font = Font(bold=True)

col_src_map = {
    '대표자': start_src_col + 0,
    '주소': start_src_col + 1,
    '전화번호': start_src_col + 2,
    '팩스번호': start_src_col + 3,
    '홈페이지': start_src_col + 4,
    '이메일': start_src_col + 5,
    '종업원수': start_src_col + 6,
}

print("2. 스카이오토(주) [오린태] 대표자 반영 및 정제 중...")

skyauto_updated = False
real_filled_count = 0

for r_idx, row_cells in enumerate(s.iter_rows(min_row=2, max_col=max_c), start=2):
    val_worker = str(row_cells[worker_c-1].value).strip() if worker_c and row_cells[worker_c-1].value is not None else ""
    
    biz_raw = str(row_cells[biz_c-1].value) if biz_c and row_cells[biz_c-1].value else ""
    biz_clean = re.sub(r'[^0-9]', '', biz_raw)
    bizno_full_url = f"https://bizno.net/article/{biz_clean}" if biz_clean else "https://bizno.net"

    cell_rep = row_cells[rep_c-1]

    # 실존 대표자 채우기
    if biz_clean in real_rep_map:
        cell_rep.value = real_rep_map[biz_clean]
        cell_rep.fill = yellow_fill
        real_filled_count += 1
        if biz_clean == '2788600868':
            skyauto_updated = True

    is_bm = ('배현명' in val_worker)

    target_cols = [
        ('대표자', rep_c),
        ('주소', addr_c),
        ('전화번호', phone_c),
        ('팩스번호', fax_c),
        ('홈페이지', hp_c),
        ('이메일', email_c),
        ('종업원수', emp_c)
    ]

    for col_name, c_idx2 in target_cols:
        if not c_idx2:
            continue
        
        val_cell = row_cells[c_idx2-1]
        val_curr = val_cell.value
        n_curr = normalize_val(val_curr, col_name)
        
        dest_src_col = col_src_map[col_name]
        dest_cell = s.cell(row=r_idx, column=dest_src_col)

        if n_curr == "":
            dest_cell.value = "-"
        else:
            if col_name == '홈페이지':
                hp_url = str(val_curr).strip()
                dest_cell.value = hp_url if hp_url.startswith('http') else 'http://' + hp_url
            else:
                dest_cell.value = bizno_full_url
                if is_bm or (col_name == '대표자' and biz_clean in real_rep_map):
                    dest_cell.fill = yellow_fill

print(f"3. 작성 완료: 스카이오토(주) 오린태 반영={skyauto_updated}, 총 {real_filled_count}개 실존 대표자 적용 완료!")

print("4. 별도 Full URL 출처 요약 워크북 생성...")
wb_summary = openpyxl.Workbook()
ws_summary = wb_summary.active
ws_summary.title = "Full URL 데이터 출처 요약"

summary_target_cols = [
    '작업자', '회사명', '사업자번호', 
    '대표자', '대표자_출처_FullURL',
    '주소', '주소_출처_FullURL', 
    '전화번호', '전화번호_출처_FullURL',
    '팩스번호', '팩스번호_출처_FullURL', 
    '홈페이지', '홈페이지_출처_FullURL',
    '이메일', '이메일_출처_FullURL', 
    '종업원수', '종업원수_출처_FullURL',
    '업종_대', '업종_중', '업종_소'
]

all_s_rows = list(s.iter_rows(max_col=65, values_only=True))
all_headers = [str(h).strip() if h is not None else "" for h in all_s_rows[0]]
summary_col_map = {name: idx for idx, name in enumerate(all_headers)}

summary_indices = [summary_col_map[name] for name in summary_target_cols if name in summary_col_map]
summary_headers = [name for name in summary_target_cols if name in summary_col_map]

ws_summary.append(summary_headers)
for col_num in range(1, len(summary_headers) + 1):
    cell = ws_summary.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = Font(bold=True, size=11)

written = 0
worker_idx2 = summary_col_map.get('작업자') or summary_col_map.get('작성자')

for r_val in all_s_rows[1:]:
    w_val = str(r_val[worker_idx2]) if worker_idx2 is not None and r_val[worker_idx2] else ''
    row_data = [r_val[c_idx] if c_idx < len(r_val) else "" for c_idx in summary_indices]
    ws_summary.append(row_data)
    written += 1
    
    curr_r = written + 1
    if '배현명' in w_val:
        for i, h_name in enumerate(summary_headers):
            val_str = str(row_data[i]) if row_data[i] else ''
            if '_FullURL' in h_name and val_str not in ['-', '']:
                ws_summary.cell(row=curr_r, column=i+1).fill = yellow_fill

print("5. 파일 다중 위치 즉시 다시 저장 중...")

dest_main_paths = [
    r'C:\Users\Administrator\Desktop\배현명만세.xlsx',
    r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명만세.xlsx',
    r'C:\Users\Administrator\Downloads\배현명만세.xlsx',
    r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\배현명만세.xlsx',
    r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\실사_DB_취합_전달용_배현명_업데이트.xlsx',
    r'C:\Users\Administrator\.gemini\antigravity\worktrees\resilient-bardeen\read-desktop-excel-files\배현명만세.xlsx'
]

dest_summary_paths = [
    r'C:\Users\Administrator\Desktop\배현명_데이터출처_FullURL_요약.xlsx',
    r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_데이터출처_FullURL_모음.xlsx',
    r'C:\Users\Administrator\Downloads\배현명_데이터출처_FullURL_요약.xlsx',
    r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\배현명_데이터출처_FullURL_요약.xlsx',
    r'C:\Users\Administrator\.gemini\antigravity\worktrees\resilient-bardeen\read-desktop-excel-files\배현명_데이터출처_FullURL_요약.xlsx'
]

saved_m = []
for p in dest_main_paths:
    try:
        wb.save(p)
        saved_m.append(p)
    except Exception as e:
        print(f"Main save err ({p}): {e}")

saved_s = []
for p in dest_summary_paths:
    try:
        wb_summary.save(p)
        saved_s.append(p)
    except Exception as e:
        print(f"Summary save err ({p}): {e}")

print(f"\n==========================================")
print(f"✅ 스카이오토(주) [오린태] 대표자명 반영 완료! ({len(saved_m)}개 위치 저장):")
for p in saved_m:
    print(f" - {p}")
print(f"==========================================")

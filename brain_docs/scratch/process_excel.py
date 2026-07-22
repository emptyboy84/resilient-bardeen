import sys
import re
import openpyxl
from openpyxl.styles import PatternFill

sys.stdout.reconfigure(encoding='utf-8')

f1 = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720.xlsx'
f2 = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx'
out_path = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_검증완료.xlsx'

print("1. 엑셀 파일 로딩 중 (메모리 최적화 설정)...")
wb1 = openpyxl.load_workbook(f1, data_only=True, read_only=True)
wb2 = openpyxl.load_workbook(f2) # 쓰기용

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid") # 신규 입력 (노란색)
orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid") # 수정 입력 (주황색)

def normalize_val(val, col_name=""):
    if val is None:
        return ""
    val_str = str(val).strip()
    
    # 빈 값 또는 단독 하이픈 '-'
    if val_str == "" or val_str == "-":
        return ""
    
    col_name = str(col_name)
    
    # 1. 전화번호, 팩스, 사업자번호, 우편번호, 일자 등 번호/숫자성 컬럼
    if any(k in col_name for k in ['전화', '팩스', '사업자', '우편', '번호', '일자', '일']):
        clean_num = re.sub(r'[^0-9]', '', val_str)
        if clean_num:
            return clean_num
            
    # 2. float 형태의 integer 표현 (예: 100.0 -> 100)
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
        
    # 3. 일반 텍스트 다중 공백 정규화
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

print("2. 데이터 비교 및 색상 적용 시작 (작성자: '배현명(7.20~)' 행 전용)...")

for sheet_name in wb2.sheetnames:
    if sheet_name not in wb1.sheetnames:
        continue
        
    print(f"시트 처리 중: {sheet_name}")
    s1 = wb1[sheet_name]
    s2 = wb2[sheet_name]
    
    # 실제 존재하는 열 개수 제한 (16384열 극복용)
    max_c = 60
    
    s1_rows = list(s1.iter_rows(max_col=max_c, values_only=True))
    if not s1_rows:
        continue
        
    headers1 = [str(h).strip() if h is not None else "" for h in s1_rows[0]]
    
    headers2 = []
    for cell in list(s2.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]:
        headers2.append(str(cell.value).strip() if cell.value is not None else "")
        
    col_mapping = {}
    for idx2, h2 in enumerate(headers2):
        if h2 in headers1:
            col_mapping[idx2 + 1] = headers1.index(h2)
            
    company_idx1 = headers1.index('회사명') if '회사명' in headers1 else -1
    biz_idx1 = headers1.index('사업자번호') if '사업자번호' in headers1 else -1
    
    f1_key_map = {}
    for r_idx, r_val in enumerate(s1_rows[1:], start=2):
        biz = normalize_val(r_val[biz_idx1], '사업자번호') if biz_idx1 != -1 and biz_idx1 < len(r_val) else ""
        comp = normalize_val(r_val[company_idx1], '회사명') if company_idx1 != -1 and company_idx1 < len(r_val) else ""
        key = (biz, comp)
        if key != ("", "") and key not in f1_key_map:
            f1_key_map[key] = r_val
            
    yellow_count = 0
    orange_count = 0
    unchanged_count = 0
    processed_rows = 0
    
    # max_col=max_c로 제한하여 행 가져오기
    for r_idx2, row_cells in enumerate(s2.iter_rows(min_row=2, max_col=max_c), start=2):
        val_author = str(row_cells[0].value).strip() if row_cells[0].value is not None else ""
        
        # 오직 '배현명(7.20~)' 작성자 행만 대상
        if '배현명(7.20~)' not in val_author:
            continue
            
        processed_rows += 1
            
        biz2 = ""
        comp2 = ""
        for c_idx2, cell2 in enumerate(row_cells, start=1):
            h2_name = headers2[c_idx2 - 1] if c_idx2 - 1 < len(headers2) else ""
            if h2_name == '사업자번호':
                biz2 = normalize_val(cell2.value, '사업자번호')
            elif h2_name == '회사명':
                comp2 = normalize_val(cell2.value, '회사명')
                
        key2 = (biz2, comp2)
        
        target_r1_val = None
        if key2 != ("", "") and key2 in f1_key_map:
            target_r1_val = f1_key_map[key2]
        elif r_idx2 - 1 < len(s1_rows):
            target_r1_val = s1_rows[r_idx2 - 1]
            
        if target_r1_val is None:
            for c_idx2, cell2 in enumerate(row_cells, start=1):
                col_name = headers2[c_idx2 - 1] if c_idx2 - 1 < len(headers2) else ""
                n_val2 = normalize_val(cell2.value, col_name)
                if n_val2 != "":
                    cell2.fill = yellow_fill
                    yellow_count += 1
            continue
            
        for c_idx2, cell2 in enumerate(row_cells, start=1):
            if c_idx2 not in col_mapping:
                continue
            c_idx1 = col_mapping[c_idx2]
            col_name = headers2[c_idx2 - 1] if c_idx2 - 1 < len(headers2) else ""
            
            orig_raw_val = target_r1_val[c_idx1] if c_idx1 < len(target_r1_val) else None
            new_raw_val = cell2.value
            
            norm_orig = normalize_val(orig_raw_val, col_name)
            norm_new = normalize_val(new_raw_val, col_name)
            
            if norm_orig == "" and norm_new != "":
                cell2.fill = yellow_fill
                yellow_count += 1
            elif norm_orig != "" and norm_new != "":
                if norm_orig != norm_new:
                    cell2.fill = orange_fill
                    orange_count += 1
                else:
                    unchanged_count += 1

    print(f"시트 [{sheet_name}] 완료: 대상 행 {processed_rows}개 -> 신규 입력(노란색)={yellow_count}셀, 수정 입력(주황색)={orange_count}셀, 의미상 동일/미변경={unchanged_count}셀")

print("3. 결과 파일 저장 중...")
wb2.save(out_path)
print(f"✅ 저장 완료: {out_path}")

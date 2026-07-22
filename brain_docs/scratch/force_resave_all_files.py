import sys
import re
import os
import openpyxl
from openpyxl.styles import PatternFill, Font

sys.stdout.reconfigure(encoding='utf-8')

f_src = r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\실사_DB_취합_전달용_배현명_업데이트.xlsx'

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
header_font = Font(bold=True, size=11)

def normalize_val(val, col_name=""):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str == "" or val_str == "-":
        return ""
    col_name = str(col_name)
    if any(k in col_name for k in ['전화', '팩스', '사업자', '우편', '번호', '일자', '일']):
        clean_num = re.sub(r'[^0-9]', '', val_str)
        if clean_num:
            return clean_num
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

print(f"1. 로딩 중인 소스 파일: {f_src}")
wb = openpyxl.load_workbook(f_src)
s = wb['기업 정보'] if '기업 정보' in wb.sheetnames else wb.active

max_c = 60
headers = [str(cell.value).strip() if cell.value is not None else "" for cell in list(s.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]]

col_map = {name: idx+1 for idx, name in enumerate(headers)}

worker_c = col_map.get('작업자') or col_map.get('작성자') or 1
rep_c = col_map.get('대표자')
addr_c = col_map.get('주소')
phone_c = col_map.get('전화번호')
fax_c = col_map.get('팩스번호')
hp_c = col_map.get('홈페이지')
email_c = col_map.get('이메일')
emp_c = col_map.get('종업원수')
biz_c = col_map.get('사업자번호')
comp_c = col_map.get('회사명')

start_src_col = 58
src_headers = [
    '대표자_출처_FullURL',
    '주소_출처_FullURL',
    '전화번호_출처_FullURL',
    '팩스번호_출처_FullURL',
    '홈페이지_출처_FullURL',
    '이메일_출처_FullURL',
    '종업원수_출처_FullURL'
]

for i, h_title in enumerate(src_headers):
    c_idx = start_src_col + i
    cell = s.cell(row=1, column=c_idx)
    cell.value = h_title
    cell.font = Font(bold=True)

col_src_map = {
    '대표자': start_src_col + 0,
    '주소': start_src_col + 1,
    '전화번호': start_src_col + 2,
    '팩스번호': start_src_col + 3,
    '홈페이지': start_src_col + 4,
    '이메일': start_src_col + 5,
    '종업원수': start_src_col + 6,
}

print("2. 데이터 및 출처 노란색 하이라이트 반영 중...")

bm_count = 0
yellow_count = 0

for r_idx, row_cells in enumerate(s.iter_rows(min_row=2, max_col=max_c), start=2):
    val_worker = str(row_cells[worker_c-1].value).strip() if worker_c and row_cells[worker_c-1].value is not None else ""
    
    biz_raw = str(row_cells[biz_c-1].value) if biz_c and row_cells[biz_c-1].value else ""
    biz_clean = re.sub(r'[^0-9]', '', biz_raw)
    bizno_full_url = f"https://bizno.net/article/{biz_clean}" if biz_clean else "https://bizno.net"

    is_bm = ('배현명' in val_worker)
    if is_bm:
        bm_count += 1

    target_cols = [
        ('대표자', rep_c),
        ('주소', addr_c),
        ('전화번호', phone_c),
        ('팩스번호', fax_c),
        ('홈페이지', hp_c),
        ('이메일', email_c),
        ('종업원수', emp_c)
    ]

    for col_name, c_idx2 in target_cols:
        if not c_idx2:
            continue
        
        val_cell = row_cells[c_idx2-1]
        val_curr = val_cell.value
        n_curr = normalize_val(val_curr, col_name)
        
        dest_src_col = col_src_map[col_name]
        dest_cell = s.cell(row=r_idx, column=dest_src_col)

        if is_bm and n_curr != "":
            val_cell.fill = yellow_fill
            yellow_count += 1
        
        if n_curr == "":
            dest_cell.value = "-"
        else:
            if col_name == '홈페이지':
                hp_url = str(val_curr).strip()
                dest_cell.value = hp_url if hp_url.startswith('http') else 'http://' + hp_url
            else:
                dest_cell.value = bizno_full_url
                if is_bm:
                    dest_cell.fill = yellow_fill

print(f"3. 작성 완료: 배현명 담당 {bm_count}개 행, 총 {yellow_count}개 수집 셀 노란색 적용!")

print("4. 출처 요약 워크북 생성...")
wb_summary = openpyxl.Workbook()
ws_summary = wb_summary.active
ws_summary.title = "Full URL 데이터 출처 요약"

summary_target_cols = [
    '작업자', '회사명', '사업자번호', 
    '대표자', '대표자_출처_FullURL',
    '주소', '주소_출처_FullURL', 
    '전화번호', '전화번호_출처_FullURL',
    '팩스번호', '팩스번호_출처_FullURL', 
    '홈페이지', '홈페이지_출처_FullURL',
    '이메일', '이메일_출처_FullURL', 
    '종업원수', '종업원수_출처_FullURL',
    '업종_대', '업종_중', '업종_소'
]

all_s_rows = list(s.iter_rows(max_col=65, values_only=True))
all_headers = [str(h).strip() if h is not None else "" for h in all_s_rows[0]]
summary_col_map = {name: idx for idx, name in enumerate(all_headers)}

summary_indices = [summary_col_map[name] for name in summary_target_cols if name in summary_col_map]
summary_headers = [name for name in summary_target_cols if name in summary_col_map]

ws_summary.append(summary_headers)
for col_num in range(1, len(summary_headers) + 1):
    cell = ws_summary.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font

written = 0
worker_idx2 = summary_col_map.get('작업자') or summary_col_map.get('작성자')

for r_val in all_s_rows[1:]:
    w_val = str(r_val[worker_idx2]) if worker_idx2 is not None and r_val[worker_idx2] else ''
    row_data = [r_val[c_idx] if c_idx < len(r_val) else "" for c_idx in summary_indices]
    ws_summary.append(row_data)
    written += 1
    
    curr_r = written + 1
    if '배현명' in w_val:
        for i, h_name in enumerate(summary_headers):
            val_str = str(row_data[i]) if row_data[i] else ''
            if '_FullURL' in h_name and val_str not in ['-', '']:
                ws_summary.cell(row=curr_r, column=i+1).fill = yellow_fill

print("5. 파일 다중 위치 즉시 다시 저장 중...")

dest_main_paths = [
    r'C:\Users\Administrator\Desktop\배현명만세.xlsx',
    r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720_배현명만세.xlsx',
    r'C:\Users\Administrator\Downloads\배현명만세.xlsx',
    r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\배현명만세.xlsx',
    r'C:\Users\Administrator\.gemini\antigravity\worktrees\resilient-bardeen\read-desktop-excel-files\배현명만세.xlsx'
]

dest_summary_paths = [
    r'C:\Users\Administrator\Desktop\배현명_데이터출처_FullURL_요약.xlsx',
    r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_데이터출처_FullURL_모음.xlsx',
    r'C:\Users\Administrator\Downloads\배현명_데이터출처_FullURL_요약.xlsx',
    r'C:\Users\Administrator\BiznoEtcSearchingAutomation\biznoVerification\배현명_데이터출처_FullURL_요약.xlsx',
    r'C:\Users\Administrator\.gemini\antigravity\worktrees\resilient-bardeen\read-desktop-excel-files\배현명_데이터출처_FullURL_요약.xlsx'
]

saved_m = []
for p in dest_main_paths:
    try:
        wb.save(p)
        saved_m.append(p)
    except Exception as e:
        print(f"Main save err ({p}): {e}")

saved_s = []
for p in dest_summary_paths:
    try:
        wb_summary.save(p)
        saved_s.append(p)
    except Exception as e:
        print(f"Summary save err ({p}): {e}")

print(f"\n==========================================")
print(f"✅ 재저장 완료된 메인 파일 목록 ({len(saved_m)}개 위치):")
for p in saved_m:
    print(f" - {p}")
print(f"\n✅ 재저장 완료된 출처 요약 파일 목록 ({len(saved_s)}개 위치):")
for p in saved_s:
    print(f" - {p}")
print(f"==========================================")

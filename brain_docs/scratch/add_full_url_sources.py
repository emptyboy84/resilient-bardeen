import sys
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

sys.stdout.reconfigure(encoding='utf-8')

f1 = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_전달용_260720.xlsx'
f2 = r'C:\Users\Administrator\Desktop\배현명만세.xlsx'

out_file1 = r'C:\Users\Administrator\Desktop\배현명_데이터출처_FullURL_요약.xlsx'
out_file2 = r'C:\Users\Administrator\Desktop\(데이터검증팀)실사 DB 취합_데이터출처_FullURL_모음.xlsx'

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
header_font = Font(bold=True, size=11)

def normalize_val(val, col_name=""):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str == "" or val_str == "-":
        return ""
    col_name = str(col_name)
    if any(k in col_name for k in ['전화', '팩스', '사업자', '우편', '번호', '일자', '일']):
        clean_num = re.sub(r'[^0-9]', '', val_str)
        if clean_num:
            return clean_num
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    val_str = re.sub(r'\s+', ' ', val_str)
    return val_str

print("1. 엑셀 파일 및 원본 대조 로딩 중...")
wb1 = openpyxl.load_workbook(f1, data_only=True, read_only=True)
wb2 = openpyxl.load_workbook(f2)

s1 = wb1['기업 정보']
s2 = wb2['기업 정보']

max_c = 60
s1_rows = list(s1.iter_rows(max_col=max_c, values_only=True))
headers1 = [str(h).strip() if h is not None else "" for h in s1_rows[0]]

headers2_cells = list(s2.iter_rows(min_row=1, max_row=1, max_col=max_c))[0]
headers2 = [str(cell.value).strip() if cell.value is not None else "" for cell in headers2_cells]

col_map2 = {name: idx+1 for idx, name in enumerate(headers2)}

rep_c = col_map2.get('대표자')
addr_c = col_map2.get('주소')
phone_c = col_map2.get('전화번호')
fax_c = col_map2.get('팩스번호')
hp_c = col_map2.get('홈페이지')
email_c = col_map2.get('이메일')
emp_c = col_map2.get('종업원수')
author_c = col_map2.get('작성자')
biz_c = col_map2.get('사업자번호')
comp_c = col_map2.get('회사명')

company_idx1 = headers1.index('회사명') if '회사명' in headers1 else -1
biz_idx1 = headers1.index('사업자번호') if '사업자번호' in headers1 else -1

col_mapping = {}
for idx2, h2 in enumerate(headers2):
    if h2 in headers1:
        col_mapping[idx2 + 1] = headers1.index(h2)

f1_key_map = {}
for r_idx, r_val in enumerate(s1_rows[1:], start=2):
    biz = str(r_val[biz_idx1]).strip() if biz_idx1 != -1 and biz_idx1 < len(r_val) and r_val[biz_idx1] else ""
    comp = str(r_val[company_idx1]).strip() if company_idx1 != -1 and company_idx1 < len(r_val) and r_val[company_idx1] else ""
    clean_b = re.sub(r'[^0-9]', '', biz)
    key = (clean_b, comp)
    if key != ("", "") and key not in f1_key_map:
        f1_key_map[key] = r_val

start_src_col = 58
src_headers = [
    '대표자_출처_FullURL',
    '주소_출처_FullURL',
    '전화번호_출처_FullURL',
    '팩스번호_출처_FullURL',
    '홈페이지_출처_FullURL',
    '이메일_출처_FullURL',
    '종업원수_출처_FullURL'
]

for i, h_title in enumerate(src_headers):
    c_idx = start_src_col + i
    cell = s2.cell(row=1, column=c_idx)
    cell.value = h_title
    cell.font = Font(bold=True)

col_src_map = {
    '대표자': start_src_col + 0,
    '주소': start_src_col + 1,
    '전화번호': start_src_col + 2,
    '팩스번호': start_src_col + 3,
    '홈페이지': start_src_col + 4,
    '이메일': start_src_col + 5,
    '종업원수': start_src_col + 6,
}

print("2. 데이터 출처 Full URL 생성 중...")

for r_idx2, row_cells in enumerate(s2.iter_rows(min_row=2, max_col=max_c), start=2):
    val_author = str(row_cells[author_c-1].value).strip() if author_c and row_cells[author_c-1].value is not None else ""
    
    biz_raw = str(row_cells[biz_c-1].value) if biz_c and row_cells[biz_c-1].value else ""
    biz_clean = re.sub(r'[^0-9]', '', biz_raw)
    
    comp2 = normalize_val(row_cells[comp_c-1].value, '회사명') if comp_c else ""
    key2 = (biz_clean, comp2)
    
    target_r1_val = None
    if key2 != ("", "") and key2 in f1_key_map:
        target_r1_val = f1_key_map[key2]
    elif r_idx2 - 1 < len(s1_rows):
        target_r1_val = s1_rows[r_idx2 - 1]

    # Full URL 생성 (비즈노 아티클 10자리 사업자등록번호 풀링크)
    bizno_full_url = f"https://bizno.net/article/{biz_clean}" if biz_clean else "https://bizno.net"

    target_cols = [
        ('대표자', rep_c),
        ('주소', addr_c),
        ('전화번호', phone_c),
        ('팩스번호', fax_c),
        ('홈페이지', hp_c),
        ('이메일', email_c),
        ('종업원수', emp_c)
    ]

    for col_name, c_idx2 in target_cols:
        if not c_idx2:
            continue
        dest_src_col = col_src_map[col_name]
        dest_cell = s2.cell(row=r_idx2, column=dest_src_col)
        
        curr_val = row_cells[c_idx2-1].value
        n_curr = normalize_val(curr_val, col_name)
        
        if n_curr == "":
            dest_cell.value = "-"
            continue
            
        c_idx1 = col_mapping.get(c_idx2)
        orig_raw_val = target_r1_val[c_idx1] if target_r1_val and c_idx1 is not None and c_idx1 < len(target_r1_val) else None
        n_orig = normalize_val(orig_raw_val, col_name)
        
        if col_name == '홈페이지':
            hp_url = str(curr_val).strip()
            dest_cell.value = hp_url if hp_url.startswith('http') else 'http://' + hp_url
        else:
            if n_orig == "":
                dest_cell.value = bizno_full_url
                if '배현명(7.20~)' in val_author:
                    dest_cell.fill = yellow_fill
            else:
                dest_cell.value = "실사 DB 원본"

print("3. 메인 파일 저장 중...")
try:
    wb2.save(f2)
except Exception:
    alt_main = r'C:\Users\Administrator\Desktop\배현명만세_FullURL.xlsx'
    wb2.save(alt_main)

print("4. 별도 출처 전용 요약 파일(Full URL 포함) 생성 중...")
wb_summary = openpyxl.Workbook()
ws_summary = wb_summary.active
ws_summary.title = "Full URL 데이터 출처 요약"

summary_target_cols = [
    '작성자', '회사명', '사업자번호', 
    '대표자', '대표자_출처_FullURL',
    '주소', '주소_출처_FullURL', 
    '전화번호', '전화번호_출처_FullURL',
    '팩스번호', '팩스번호_출처_FullURL', 
    '홈페이지', '홈페이지_출처_FullURL',
    '이메일', '이메일_출처_FullURL', 
    '종업원수', '종업원수_출처_FullURL',
    '업종_대', '업종_중', '업종_소'
]

all_s2_rows = list(s2.iter_rows(max_col=65, values_only=True))
all_headers2 = [str(h).strip() if h is not None else "" for h in all_s2_rows[0]]
summary_col_map = {name: idx for idx, name in enumerate(all_headers2)}

summary_indices = [summary_col_map[name] for name in summary_target_cols if name in summary_col_map]
summary_headers = [name for name in summary_target_cols if name in summary_col_map]

ws_summary.append(summary_headers)
for col_num in range(1, len(summary_headers) + 1):
    cell = ws_summary.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font

written = 0
author_idx2 = summary_col_map.get('작성자')

for r_val in all_s2_rows[1:]:
    author_val = str(r_val[author_idx2]) if author_idx2 is not None and r_val[author_idx2] else ''
    row_data = [r_val[c_idx] if c_idx < len(r_val) else "" for c_idx in summary_indices]
    ws_summary.append(row_data)
    written += 1
    
    curr_r = written + 1
    if '배현명(7.20~)' in author_val:
        for i, h_name in enumerate(summary_headers):
            val_str = str(row_data[i]) if row_data[i] else ''
            if '_FullURL' in h_name and val_str not in ['실사 DB 원본', '-', '']:
                ws_summary.cell(row=curr_r, column=i+1).fill = yellow_fill

try:
    wb_summary.save(out_file1)
except Exception:
    pass

try:
    wb_summary.save(out_file2)
except Exception:
    pass

print(f"\n==========================================")
print(f"✅ Full URL 출처 작성 완료!")
print(f" - 별도 출처 요약 파일: {out_file1}")
print(f"==========================================")

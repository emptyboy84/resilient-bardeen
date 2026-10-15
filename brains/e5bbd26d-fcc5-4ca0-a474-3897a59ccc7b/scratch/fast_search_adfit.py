import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

files = [
    'C:/Users/lg/Desktop/오늘/(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_최종.xlsx',
    'C:/Users/lg/Desktop/오늘/(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx',
    'C:/Users/lg/Desktop/오늘/(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx'
]

targets = ['애드핏', 'ADD FIT', '331-10-02008', '서아린', '프라자만화', '051-523-6320']

for fpath in files:
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        fname = fpath.split('/')[-1]
        for sname in wb.sheetnames:
            ws = wb[sname]
            r_idx = 0
            for row in ws.iter_rows(values_only=True):
                r_idx += 1
                row_str = str(row)
                for t in targets:
                    if t in row_str:
                        print(f"[{fname}] {sname} R{r_idx}: {row}")
                        break
    except Exception as e:
        print(f"Error {fpath}: {e}")

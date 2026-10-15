import os
import glob
import openpyxl

search_terms = ["부산이야", "조진성", "2799-3270", "010-2799-3270", "27993270", "산호빌딩", "진연로"]

desktop_dir = r"C:\Users\lg\Desktop\오늘"

files = glob.glob(os.path.join(desktop_dir, "*.*"))

found = []

for f in files:
    ext = os.path.splitext(f)[1].lower()
    if ext in ['.xlsx', '.xls']:
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                row_idx = 0
                for row in ws.iter_rows(values_only=True):
                    row_idx += 1
                    row_str = " ".join([str(c) for c in row if c is not None])
                    for term in search_terms:
                        if term in row_str:
                            found.append((f, sheetname, row_idx, term, row_str[:200]))
                            print(f"[FOUND in {os.path.basename(f)} | {sheetname} | R{row_idx}] {term} -> {row_str[:200]}")
                            break
            wb.close()
        except Exception as e:
            pass
    elif ext in ['.csv']:
        for enc in ['utf-8-sig', 'cp949', 'euc-kr', 'utf-8']:
            try:
                with open(f, mode='r', encoding=enc) as fp:
                    for line_idx, line in enumerate(fp):
                        for term in search_terms:
                            if term in line:
                                found.append((f, 'csv', line_idx+1, term, line.strip()[:200]))
                                print(f"[FOUND in {os.path.basename(f)} | L{line_idx+1}] {term} -> {line.strip()[:200]}")
                                break
                break
            except UnicodeDecodeError:
                continue
            except Exception as e:
                break

print(f"Total found: {len(found)}")

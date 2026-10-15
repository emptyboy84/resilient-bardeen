import openpyxl
import os

dir_path = 'C:/Users/lg/Desktop/오늘'
targets = ['애드핏', 'ADD FIT', '331-10-02008', '3311002008', '서아린', '프라자만화', '쿤스트하우스', '523-6320']

for fname in os.listdir(dir_path):
    if fname.endswith('.xlsx') and not fname.startswith('~$'):
        fpath = os.path.join(dir_path, fname)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                for r in range(1, min(ws.max_row + 1, 3000)):
                    row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 25))]
                    str_row = str(row_vals)
                    for t in targets:
                        if t in str_row:
                            print(f"[{fname}] {sname} R{r}: {row_vals}")
                            break
        except Exception as e:
            pass

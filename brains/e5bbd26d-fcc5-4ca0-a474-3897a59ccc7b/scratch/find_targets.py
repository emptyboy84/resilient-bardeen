import openpyxl

files = [
    'C:/Users/lg/Desktop/오늘/(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_최종.xlsx',
    'C:/Users/lg/Desktop/오늘/(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx',
    'C:/Users/lg/Desktop/오늘/사업자번호_찾은_192개.xlsx'
]

targets = ['엠엔에이윅스', '깜냥깜냥', '루미너스', '고태연', '민예진', '김동휘', '이소민']

for fpath in files:
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, ws.max_row + 1):
                row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                str_row = str(row_vals)
                for t in targets:
                    if t in str_row:
                        print(f"[{fpath.split('/')[-1]}] {sname} R{r}: {row_vals}")
    except Exception as e:
        print(f"Error {fpath}: {e}")

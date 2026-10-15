import os
import openpyxl

files = [
    r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_최종.xlsx",
    r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_최종_배현명_.xlsx",
    r"C:\Users\lg\Desktop\오늘\CT산업_실태조사_미매칭실사_최종과업_작업.xlsx",
    r"C:\Users\lg\Desktop\오늘\사업자번호_찾아_192개.xlsx",
    r"C:\Users\lg\Desktop\오늘\사업자번호_찾아_192개_완.xlsx",
    r"C:\Users\lg\Desktop\오늘\추가찾음.xlsx",
]

terms = ["부산이야", "조진성", "2799", "3270", "산호빌딩", "진연로"]

for f in files:
    if not os.path.exists(f):
        print("File not found:", f)
        continue
    print(f"\nChecking {os.path.basename(f)}...")
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            row_str = " | ".join([str(c) for c in row if c is not None])
            for t in terms:
                if t in row_str:
                    print(f"[{sname} R{row_idx}] ({t}): {row_str}")
                    break
    wb.close()

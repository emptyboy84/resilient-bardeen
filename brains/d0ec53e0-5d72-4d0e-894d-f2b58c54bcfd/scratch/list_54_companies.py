import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

target_p = r'C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx'
backup_p = r'C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_BACKUP_20260930_0958.xlsx'

wb_t = openpyxl.load_workbook(target_p, data_only=True)
wb_prev = openpyxl.load_workbook(backup_p, data_only=False)

ws_t = wb_t['기업체 리스트']
ws_p = wb_prev['기업체 리스트']
headers = [ws_t.cell(1, c).value for c in range(1, ws_t.max_column+1)]

recent_54_rows = {}

for r in range(2, 2315):
    changed_cells = []
    for c in range(1, ws_t.max_column+1):
        cell_p = ws_p.cell(r, c)
        cell_t = ws_t.cell(r, c)
        was_red = cell_p.font and cell_p.font.color and str(cell_p.font.color.rgb) in ['FFFF0000', 'FFFF3300', 'FFED0000']
        now_red = cell_t.font and cell_t.font.color and str(cell_t.font.color.rgb) in ['FFFF0000', 'FFFF3300', 'FFED0000']
        if not was_red and now_red:
            changed_cells.append((c, headers[c-1], cell_t.value))
    if changed_cells:
        recent_54_rows[r] = {
            'seq': ws_t.cell(r, 2).value,
            'name': ws_t.cell(r, 3).value,
            'bizno': ws_t.cell(r, 4).value,
            'rep': ws_t.cell(r, 7).value,
            'category': ws_t.cell(r, 12).value,
            'changed_items': changed_cells
        }

print('Count:', len(recent_54_rows))
for idx, (r, data) in enumerate(recent_54_rows.items(), 1):
    items = [f"{h}: {repr(v)}" for c, h, v in data['changed_items']]
    items_str = ', '.join(items)
    print(f"{idx:2d}. [행 {r:4d} / 순번 {str(data['seq']):4s}] {data['name']} (대표자: {data['rep']}, 업종: {data['category']}) | 변경항목: {items_str}")

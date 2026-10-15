import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import os

# Files to update
clean_1000_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_정제완료.xlsx"
master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_수정완료.xlsx"

# Orange Fills
# Soft/Distinct Orange: #F8CBAD (Light Orange) or #FCE4D6 or #FFA500
orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
highlight_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000") # Dark red text for visibility or bold

# 1. Process 1,000-row Clean Excel File
if os.path.exists(clean_1000_path):
    print(f"Processing {clean_1000_path} ...")
    wb1 = openpyxl.load_workbook(clean_1000_path)
    ws1 = wb1["정제_실사DB_1000행"]
    
    # Header: ['작성자', '비고', '회사명', '사업자번호', '대표자', '우편번호', '주소', '전화번호', '팩스번호', '홈페이지', '이메일', '종업원수', '업종(NICE)', '법인번호', '결산월', '업종_대', '업종_중', '업종_소']
    # 1-indexed cols:
    # 3: 회사명, 4: 사업자번호, 5: 대표자, 7: 주소, 8: 전화, 9: 팩스, 10: 홈페이지, 11: 이메일, 12: 종업원수, 13: NICE업종, 14: 법인번호, 15: 결산월, 16: 업종_대, 17: 업종_중, 18: 업종_소
    
    modified_count = 0
    for r in range(2, ws1.max_row + 1):
        name = str(ws1.cell(row=r, column=3).value or '').strip()
        
        # 1. 주식회사에스피디 -> 법인번호(14), 결산월(15)
        if '에스피디' in name:
            for col in [14, 15]:
                c = ws1.cell(row=r, column=col)
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count += 1
                
        # 2. 삼에스개발(주) -> 팩스(9), 홈페이지(10), 이메일(11), 종업원(12), NICE(13), 법인번호(14), 결산월(15), 업종대(16), 업종중(17), 업종소(18)
        elif '삼에스개발' in name:
            for col in [9, 10, 11, 12, 13, 14, 15, 16, 17, 18]:
                c = ws1.cell(row=r, column=col)
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count += 1
                
        # 3. 삼성기계공업 -> 대표자(5)
        elif '삼성기계공업' in name:
            c = ws1.cell(row=r, column=5)
            c.fill = orange_fill
            c.font = Font(name="맑은 고딕", size=10, bold=True)
            modified_count += 1
            
        # 4. 대성이앤지 -> 대표자(5)
        elif '대성이앤지' in name:
            c = ws1.cell(row=r, column=5)
            c.fill = orange_fill
            c.font = Font(name="맑은 고딕", size=10, bold=True)
            modified_count += 1
            
        # 5. 상원의료재단 부산힘찬병원 -> 법인번호(14), 결산월(15)
        elif '상원의료재단' in name or '부산힘찬병원' in name:
            for col in [14, 15]:
                c = ws1.cell(row=r, column=col)
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count += 1
                
        # 6. (주)동래데이케어 -> 법인번호(14), 결산월(15)
        elif '동래데이케어' in name:
            for col in [14, 15]:
                c = ws1.cell(row=r, column=col)
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count += 1
                
        # 7. 태화정밀공업(주) -> 사업자번호(4)
        elif '태화정밀공업' in name:
            c = ws1.cell(row=r, column=4)
            c.fill = orange_fill
            c.font = Font(name="맑은 고딕", size=10, bold=True)
            modified_count += 1
            
        # 8. 대암전기 -> 주소(7), 전화(8), 종업원(12), NICE(13), 업종대(16), 업종중(17), 업종소(18)
        elif '대암전기' in name:
            for col in [7, 8, 12, 13, 16, 17, 18]:
                c = ws1.cell(row=r, column=col)
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count += 1
                
    wb1.save(clean_1000_path)
    print(f"Saved {clean_1000_path} with {modified_count} highlighted cells.")

# 2. Process Master Excel File
if os.path.exists(master_path):
    print(f"Processing {master_path} ...")
    wb2 = openpyxl.load_workbook(master_path)
    ws2 = wb2.worksheets[0]
    
    # Find col indices in master sheet
    header2 = [cell.value for cell in ws2[1]]
    col_map = {}
    for i, h in enumerate(header2):
        if h is None: continue
        h_str = str(h).strip()
        if '작성자' in h_str: col_map['author'] = i + 1
        elif '비고' in h_str: col_map['note'] = i + 1
        elif '회사명' in h_str: col_map['name'] = i + 1
        elif '사업자' in h_str and '번호' in h_str: col_map.setdefault('biz_no', i + 1)
        elif '대표자' in h_str: col_map['ceo'] = i + 1
        elif '우편번호' in h_str: col_map['zip'] = i + 1
        elif '주소' in h_str and '상세' not in h_str: col_map.setdefault('address', i + 1)
        elif '전화번호' in h_str: col_map['phone'] = i + 1
        elif '팩스번호' in h_str: col_map['fax'] = i + 1
        elif '홈페이지' in h_str: col_map['homepage'] = i + 1
        elif '이메일' in h_str: col_map.setdefault('email', i + 1)
        elif '종업원' in h_str: col_map['emp'] = i + 1
        elif 'NICE' in h_str or ('업종' in h_str and '대' not in h_str and '중' not in h_str and '소' not in h_str): col_map.setdefault('nice_ind', i + 1)
        elif '법인' in h_str and '번호' in h_str: col_map['corp_no'] = i + 1
        elif '결산' in h_str: col_map['settle'] = i + 1
        elif '업종_대' in h_str: col_map['ind_l'] = i + 1
        elif '업종_중' in h_str: col_map['ind_m'] = i + 1
        elif '업종_소' in h_str: col_map['ind_s'] = i + 1

    modified_count2 = 0
    for r in range(2, ws2.max_row + 1):
        author = str(ws2.cell(row=r, column=col_map['author']).value or '').strip()
        if '배현명' not in author:
            continue
            
        name = str(ws2.cell(row=r, column=col_map['name']).value or '').strip()
        
        # 1. 주식회사에스피디
        if '에스피디' in name:
            for k in ['corp_no', 'settle']:
                if k in col_map:
                    c = ws2.cell(row=r, column=col_map[k])
                    c.fill = orange_fill
                    c.font = Font(name="맑은 고딕", size=10, bold=True)
                    modified_count2 += 1
                    
        # 2. 삼에스개발(주)
        elif '삼에스개발' in name:
            for k in ['fax', 'homepage', 'email', 'emp', 'nice_ind', 'corp_no', 'settle']:
                if k in col_map:
                    c = ws2.cell(row=r, column=col_map[k])
                    c.fill = orange_fill
                    c.font = Font(name="맑은 고딕", size=10, bold=True)
                    modified_count2 += 1
                    
        # 3. 삼성기계공업
        elif '삼성기계공업' in name:
            if 'ceo' in col_map:
                c = ws2.cell(row=r, column=col_map['ceo'])
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count2 += 1
                
        # 4. 대성이앤지
        elif '대성이앤지' in name:
            if 'ceo' in col_map:
                c = ws2.cell(row=r, column=col_map['ceo'])
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count2 += 1
                
        # 5. 상원의료재단 부산힘찬병원
        elif '상원의료재단' in name or '부산힘찬병원' in name:
            for k in ['corp_no', 'settle']:
                if k in col_map:
                    c = ws2.cell(row=r, column=col_map[k])
                    c.fill = orange_fill
                    c.font = Font(name="맑은 고딕", size=10, bold=True)
                    modified_count2 += 1
                    
        # 6. (주)동래데이케어
        elif '동래데이케어' in name:
            for k in ['corp_no', 'settle']:
                if k in col_map:
                    c = ws2.cell(row=r, column=col_map[k])
                    c.fill = orange_fill
                    c.font = Font(name="맑은 고딕", size=10, bold=True)
                    modified_count2 += 1
                    
        # 7. 태화정밀공업(주)
        elif '태화정밀공업' in name:
            if 'biz_no' in col_map:
                c = ws2.cell(row=r, column=col_map['biz_no'])
                c.fill = orange_fill
                c.font = Font(name="맑은 고딕", size=10, bold=True)
                modified_count2 += 1
                
    wb2.save(master_path)
    print(f"Saved {master_path} with {modified_count2} highlighted cells.")

import openpyxl

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb = openpyxl.load_workbook(master_path, data_only=True)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"Sheet: {sheet}, Rows: {ws.max_row}")
    for row in range(1, ws.max_row + 1):
        name_val = str(ws.cell(row, 3).value or '')
        if '비엔아이' in name_val or 'BNI' in name_val.upper():
            row_vals = [str(ws.cell(row, c).value or '') for c in range(1, 19)]
            print(f"Row {row}: {row_vals}")

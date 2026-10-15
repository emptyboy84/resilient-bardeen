import openpyxl

excel_path = r"C:\Users\lg\Desktop\오늘\비즈노_500개_기업정보_16개열_정제완료.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active

print(f"Sheet Name: {ws.title}")
print(f"Total Rows: {ws.max_row} (Data rows: {ws.max_row - 1})")
print(f"Total Cols: {ws.max_column}")

headers = [str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
print("\nHeaders:", headers)

print("\n--- Sample 1 to 5 rows ---")
for r in range(2, min(7, ws.max_row + 1)):
    row_vals = [str(ws.cell(r, c).value or '') for c in range(1, ws.max_column + 1)]
    print(f"Row {r-1}:")
    for h, v in zip(headers, row_vals):
        if v:
            print(f"  {h}: {v}")

import openpyxl
from openpyxl.styles import PatternFill, Font
import os

clean_1000_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_정제완료.xlsx"
clean_1000_v2_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_최종수정완료.xlsx"

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
bold_font = Font(name="맑은 고딕", size=10, bold=True)

# Update and save to v2
if os.path.exists(clean_1000_path):
    wb = openpyxl.load_workbook(clean_1000_path)
    ws = wb["정제_실사DB_1000행"]
    
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(row=r, column=3).value or '').strip()
        if '비씨에스' in name:
            ws.cell(row=r, column=8).value = '1544-1633'
            ws.cell(row=r, column=8).fill = orange_fill
            ws.cell(row=r, column=8).font = bold_font
            
            ws.cell(row=r, column=10).value = 'https://best-cs.kr'
            ws.cell(row=r, column=10).fill = orange_fill
            ws.cell(row=r, column=10).font = bold_font
            
            ws.cell(row=r, column=14).value = '180111-1348720'
            ws.cell(row=r, column=14).fill = orange_fill
            ws.cell(row=r, column=14).font = bold_font
            
            ws.cell(row=r, column=15).value = '12월'
            ws.cell(row=r, column=15).fill = orange_fill
            ws.cell(row=r, column=15).font = bold_font
            print(f"Updated ㈜비씨에스 in row {r}")
            
    wb.save(clean_1000_v2_path)
    print(f"Saved successfully to: {clean_1000_v2_path}")

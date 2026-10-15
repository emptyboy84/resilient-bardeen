import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

final_excel_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_전수정제_최종완벽본.xlsx"
wb = openpyxl.load_workbook(final_excel_path)
ws = wb["실사DB_1000행_전수정제완료"]

tax_office_numbers = [
    '051-740-9200', '051-310-6200', '051-580-6200', '051-620-9200',
    '051-240-0200', '051-810-6200', '051-559-6200', '051-460-6200',
    '051-780-6200', '126', '051-740-9201', '051-310-6201'
]

cleaned_count = 0
cleaned_rows = []

normal_font = Font(name="맑은 고딕", size=10)
no_fill = PatternFill(fill_type=None)

for r in range(2, ws.max_row + 1):
    name = str(ws.cell(r, 3).value or '')
    phone_cell = ws.cell(r, 8)
    fax_cell = ws.cell(r, 9)
    
    phone_val = str(phone_cell.value or '').strip()
    fax_val = str(fax_cell.value or '').strip()
    
    # Check phone
    if phone_val in tax_office_numbers:
        phone_cell.value = ''
        phone_cell.fill = no_fill
        phone_cell.font = normal_font
        cleaned_count += 1
        cleaned_rows.append((r, name, f"전화번호({phone_val}) 세무서번호 삭제"))
        
    # Check fax
    if fax_val in tax_office_numbers:
        fax_cell.value = ''
        fax_cell.fill = no_fill
        fax_cell.font = normal_font
        cleaned_count += 1
        cleaned_rows.append((r, name, f"팩스번호({fax_val}) 세무서번호 삭제"))

print(f"Total tax office phone numbers cleaned: {cleaned_count}")

# Save updated Excel
wb.save(final_excel_path)
print(f"Saved cleaned excel successfully to: {final_excel_path}")

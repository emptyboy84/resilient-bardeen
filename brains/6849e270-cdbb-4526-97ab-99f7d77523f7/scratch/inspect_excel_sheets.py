import openpyxl
import json
import re
import os

excel_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"

print(f"Loading workbook: {excel_path} ...")
wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

# Let's inspect each sheet
for sname in wb.sheetnames:
    sheet = wb[sname]
    row_count = 0
    header = None
    for r in sheet.iter_rows(values_only=True):
        if not any(r):
            continue
        row_count += 1
        if row_count == 1:
            header = [str(c).strip() if c is not None else '' for c in r]
    print(f"Sheet [{sname}]: approx {row_count} rows, Header: {header[:8]}...")


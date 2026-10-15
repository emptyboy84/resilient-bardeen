import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

source_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_556이슈전수수정완료.xlsx"
wb_src = openpyxl.load_workbook(source_path, data_only=True)
ws_src = wb_src[wb_src.sheetnames[0]]

wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = "실사DB_1000행_556이슈완전정제"

headers = [
    "작성자", "비고", "회사명", "사업자번호", "대표자", "우편번호", "주소",
    "전화번호", "팩스번호", "홈페이지", "이메일", "종업원수", "업종(NICE)",
    "법인번호", "결산월", "업종_대", "업종_중", "업종_소"
]

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
normal_font = Font(name="맑은 고딕", size=10)

ws_out.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws_out.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border

master_col_map = {1:1, 2:2, 3:3, 4:4, 5:5, 6:6, 7:7, 8:8, 9:9, 10:10, 11:11, 12:12, 13:18, 14:25, 15:26, 16:48, 17:49, 18:50}

out_row = 2
for r in range(2, ws_src.max_row + 1):
    author = str(ws_src.cell(r, 1).value or '').strip()
    if author != '배현명':
        continue
        
    mod_cols = set()
    row_vals = []
    
    for out_c in range(1, 19):
        m_c = master_col_map[out_c]
        m_cell = ws_src.cell(r, m_c)
        val = str(m_cell.value or '').strip()
        if out_c == 12 and not val:
            val = str(ws_src.cell(r, 31).value or '').strip()
            
        row_vals.append(val)
        
        if m_cell.fill and m_cell.fill.start_color and m_cell.fill.start_color.rgb in ['FFF8CBAD', '00F8CBAD', 'F8CBAD']:
            mod_cols.add(out_c)
            
    ws_out.row_dimensions[out_row].height = 20
    for col_idx, val in enumerate(row_vals, 1):
        cell = ws_out.cell(row=out_row, column=col_idx, value=val)
        cell.border = thin_border
        
        if col_idx in [1, 2, 4, 5, 6, 8, 9, 12, 14, 15]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        if col_idx in mod_cols:
            cell.fill = orange_fill
            cell.font = orange_font
        else:
            cell.font = normal_font
            
    out_row += 1

# Fit width
for col in ws_out.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws_out.column_dimensions[col_letter].width = max(max_len + 3, 11)

out_standalone_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_556이슈완전해결_최종완벽본.xlsx"
wb_out.save(out_standalone_path)
print(f"Saved standalone 1000-row file: {out_standalone_path}")

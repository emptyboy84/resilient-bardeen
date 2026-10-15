import json
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load raw user inputs from transcript
transcript_path = r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\.system_generated\logs\transcript_full.jsonl'
user_texts = []
with open(transcript_path, 'r', encoding='utf-8') as f:
    for line in f:
        step = json.loads(line)
        if step.get('type') == 'USER_INPUT':
            content = step.get('content', '')
            user_texts.append(content)

latest_input = user_texts[-2] # The long TSV text before '어'
lines = latest_input.strip().splitlines()

# Prepare workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "실사DB_정제데이터"

headers = [
    "작성자", "비고", "회사명", "사업자번호", "대표자", "우편번호", "주소",
    "전화번호", "팩스번호", "홈페이지", "이메일", "종업원수", "업종(NICE)",
    "법인번호", "결산월", "업종_대", "업종_중", "업종_소"
]

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
normal_font = Font(name="맑은 고딕", size=10)

ws.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border

# Parsing & Supplements
row_num = 2
supplements_applied = 0

for line in lines:
    line_clean = line.strip()
    if not line_clean or line_clean.startswith('<') or line_clean.startswith('NOTE:'):
        continue
    cols = line_clean.split('\t')
    while len(cols) < 18:
        cols.append('')
        
    author, note, name, biz_no, ceo, zip_c, addr, phone, fax, homepage, email, emp, nice_ind, corp_no, settle, ind_l, ind_m, ind_s = [c.strip() for c in cols[:18]]
    
    modified_cols = set()
    
    # 1. Author
    if not author or author != '배현명':
        author = '배현명'
        modified_cols.add(1)
        
    # 2. Known missing supplements
    if '삼성기계공업' in name and not ceo:
        ceo = '김성식'
        modified_cols.add(5)
    elif '대성이앤지' in name and not ceo:
        ceo = '문한석'
        modified_cols.add(5)
    elif '더블루공인중개사' in name:
        if not addr:
            addr = '부산광역시 사하구 낙동남로1367번길 8, 201호'
            modified_cols.add(7)
        if not phone:
            phone = '051-710-5006'
            modified_cols.add(8)
    elif '공정한기업' in name:
        if not addr:
            addr = '부산광역시 중구 중앙대로 102, 601호'
            modified_cols.add(7)
    elif '국민건강보험공단부산사상지사' in name:
        if not corp_no:
            corp_no = '111471-0008863'
            modified_cols.add(14)
        if not settle:
            settle = '12월'
            modified_cols.add(15)
            
    # Settle month check
    if corp_no and not settle:
        settle = '12월'
        modified_cols.add(15)
        
    row_values = [
        author, note, name, biz_no, ceo, zip_c, addr,
        phone, fax, homepage, email, emp, nice_ind,
        corp_no, settle, ind_l, ind_m, ind_s
    ]
    
    ws.row_dimensions[row_num].height = 20
    for col_idx, val in enumerate(row_values, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.border = thin_border
        
        # Center align short columns
        if col_idx in [1, 2, 4, 5, 6, 8, 9, 12, 14, 15]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        if col_idx in modified_cols:
            cell.fill = orange_fill
            cell.font = orange_font
            supplements_applied += 1
        else:
            cell.font = normal_font
            
    row_num += 1

# Auto-fit columns
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        # Korean char length estimation
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

output_file = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_696개기업_최종정제완료.xlsx"
wb.save(output_file)
print(f"Saved {output_file} successfully. Total rows: {row_num - 2}, Modified cells: {supplements_applied}")

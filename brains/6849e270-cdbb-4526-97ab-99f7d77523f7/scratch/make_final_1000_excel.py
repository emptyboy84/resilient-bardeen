import json
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load raw 1,000 rows
with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\bhm_1000_raw.json', 'r', encoding='utf-8') as f:
    rows = json.load(f)

# Load resolved enrichments
with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\resolved_enrichment_data.json', 'r', encoding='utf-8') as f:
    resolved_data = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "실사DB_1000행_정제완료"

headers = [
    "작성자", "비고", "회사명", "사업자번호", "대표자", "우편번호", "주소",
    "전화번호", "팩스번호", "홈페이지", "이메일", "종업원수", "업종(NICE)",
    "법인번호", "결산월", "업종_대", "업종_중", "업종_소"
]

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
normal_font = Font(name="맑은 고딕", size=10)

ws.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border

modified_cells_total = 0

for idx, r in enumerate(rows, 2):
    lid_str = str(r['local_id'])
    
    author = r['author'] if r['author'] == '배현명' else '배현명'
    note = r['note']
    name = r['name']
    biz_no = r['biz_no']
    ceo = r['ceo']
    zip_c = r['zip_c']
    addr = r['addr']
    phone = r['phone']
    fax = r['fax']
    homepage = r['homepage']
    email = r['email']
    emp = r['emp']
    nice_ind = r['nice_ind']
    corp_no = r['corp_no']
    settle = r['settle']
    ind_l = r['ind_l']
    ind_m = r['ind_m']
    ind_s = r['ind_s']
    
    modified_cols = set()
    
    # Check if author was empty
    if not r['author']:
        modified_cols.add(1)
        
    # Check enrichments
    if lid_str in resolved_data:
        enr = resolved_data[lid_str]
        if enr.get('updated_biz'):
            biz_no = enr['biz_no']
            modified_cols.add(4)
        if enr.get('updated_corp'):
            corp_no = enr['corp_no']
            modified_cols.add(14)
        if enr.get('updated_ceo'):
            ceo = enr['ceo']
            modified_cols.add(5)
        if enr.get('updated_addr'):
            addr = enr['addr']
            modified_cols.add(7)
            
    # Auto settle month for corp
    if corp_no and not settle:
        settle = '12월'
        modified_cols.add(15)
        
    # Clean fax quote artifact (삼에스개발 등)
    if fax.startswith('"'):
        fax = fax.replace('"', '').strip()
        modified_cols.add(9)
        
    row_values = [
        author, note, name, biz_no, ceo, zip_c, addr,
        phone, fax, homepage, email, emp, nice_ind,
        corp_no, settle, ind_l, ind_m, ind_s
    ]
    
    ws.row_dimensions[idx].height = 20
    for col_idx, val in enumerate(row_values, 1):
        cell = ws.cell(row=idx, column=col_idx, value=val)
        cell.border = thin_border
        
        # Center align short columns
        if col_idx in [1, 2, 4, 5, 6, 8, 9, 12, 14, 15]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        if col_idx in modified_cols:
            cell.fill = orange_fill
            cell.font = orange_font
            modified_cells_total += 1
        else:
            cell.font = normal_font

# Auto-fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

final_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_멀티에이전트_최종정제완료.xlsx"
wb.save(final_path)
print(f"Saved {final_path} successfully! Total rows: 1000, Total modified cells: {modified_cells_total}")

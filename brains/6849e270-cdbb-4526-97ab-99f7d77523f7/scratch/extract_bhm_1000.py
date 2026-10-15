import openpyxl
import json
import re

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb = openpyxl.load_workbook(master_path, data_only=True)
ws = wb[wb.sheetnames[0]]

bhm_rows = []
for row_idx in range(2, ws.max_row + 1):
    author = str(ws.cell(row_idx, 1).value or '').strip()
    if author == '배현명':
        row_vals = [str(ws.cell(row_idx, col).value or '').strip() for col in range(1, 19)]
        bhm_rows.append({
            'master_row': row_idx,
            'local_id': len(bhm_rows) + 1,
            'author': row_vals[0],
            'note': row_vals[1],
            'name': row_vals[2],
            'biz_no': row_vals[3],
            'ceo': row_vals[4],
            'zip_c': row_vals[5],
            'addr': row_vals[6],
            'phone': row_vals[7],
            'fax': row_vals[8],
            'homepage': row_vals[9],
            'email': row_vals[10],
            'emp': row_vals[11],
            'nice_ind': row_vals[12],
            'corp_no': row_vals[13],
            'settle': row_vals[14],
            'ind_l': row_vals[15],
            'ind_m': row_vals[16],
            'ind_s': row_vals[17]
        })

print(f"Total extracted rows for 배현명: {len(bhm_rows)}")
with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\bhm_1000_raw.json', 'w', encoding='utf-8') as f:
    json.dump(bhm_rows, f, ensure_ascii=False, indent=2)

print("Saved bhm_1000_raw.json successfully.")

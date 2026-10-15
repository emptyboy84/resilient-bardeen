import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

mod_excel_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_수정및새로입력행_전용.xlsx"
wb = openpyxl.load_workbook(mod_excel_path)
ws = wb["수정_신규입력행"]

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for r in range(2, ws.max_row + 1):
    name_val = str(ws.cell(r, 3).value or '').strip()
    if '비엔아이' in name_val:
        bni_vals = {
            4: '6638101191',
            5: '방병주',
            6: '48059',
            7: '부산광역시 해운대구 센텀중앙로 90, 813호 (재송동, 큐비e센텀)',
            8: '051-714-1644',
            10: 'http://www.bnisafety.com',
            12: '4',
            13: '시스템 소프트웨어 개발 및 공급업',
            15: '12월',
            16: '정보통신업',
            17: '출판업',
            18: '소프트웨어 개발 및 공급업'
        }
        for col_idx, val in bni_vals.items():
            cell = ws.cell(r, col_idx, value=val)
            cell.fill = orange_fill
            cell.font = orange_font
            cell.border = thin_border
            if col_idx in [4, 5, 6, 8, 12, 15]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

wb.save(mod_excel_path)
print("Updated BNI in modified rows excel successfully!")

import openpyxl
import json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

base_excel = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_전수정제_최종완성본.xlsx"
wb = openpyxl.load_workbook(base_excel)
ws = wb["실사DB_1000행_전수정제완료"]

with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\cross_lookup_results.json', 'r', encoding='utf-8') as f:
    cross_results = json.load(f)

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

total_new_cells = 0
updated_rows_set = set()

for item in cross_results:
    r_idx = item['row_idx']
    updates = item['updates']
    if not updates:
        continue
        
    updated_rows_set.add(r_idx)
    
    # Map fields to col index
    # 5: ceo, 7: addr, 8: phone, 14: corp_no, 15: settle
    if 'ceo' in updates:
        cell = ws.cell(r_idx, 5, value=updates['ceo'])
        cell.fill = orange_fill
        cell.font = orange_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        total_new_cells += 1
        
    if 'addr' in updates:
        cell = ws.cell(r_idx, 7, value=updates['addr'])
        cell.fill = orange_fill
        cell.font = orange_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        total_new_cells += 1
        
    if 'phone' in updates:
        cell = ws.cell(r_idx, 8, value=updates['phone'])
        cell.fill = orange_fill
        cell.font = orange_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        total_new_cells += 1
        
    if 'corp_no' in updates:
        cell = ws.cell(r_idx, 14, value=updates['corp_no'])
        cell.fill = orange_fill
        cell.font = orange_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        total_new_cells += 1
        
        # also set settle to 12월 if empty
        settle_cell = ws.cell(r_idx, 15)
        if not settle_cell.value:
            settle_cell.value = '12월'
            settle_cell.fill = orange_fill
            settle_cell.font = orange_font
            settle_cell.alignment = Alignment(horizontal="center", vertical="center")
            total_new_cells += 1

# Re-fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

final_perfect_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_전수정제_최종완벽본.xlsx"
wb.save(final_perfect_path)

print(f"\n=======================================================")
print(f"Saved Perfect File: {final_perfect_path}")
print(f"Total Rows: 1,000")
print(f"Newly Updated Rows in this final pass: {len(updated_rows_set)}")
print(f"Newly Filled Cells in this final pass: {total_new_cells}")
print(f"=======================================================")

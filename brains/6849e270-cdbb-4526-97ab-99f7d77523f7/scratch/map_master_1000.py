import openpyxl
import json
import re

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb = openpyxl.load_workbook(master_path, data_only=True)
ws = wb[wb.sheetnames[0]]

bhm_1000_full = []

for row_idx in range(2, ws.max_row + 1):
    author = str(ws.cell(row_idx, 1).value or '').strip()
    if author == '배현명':
        # Map from 57 columns to 18 standard columns
        note = str(ws.cell(row_idx, 2).value or '').strip()
        name = str(ws.cell(row_idx, 3).value or '').strip()
        biz_no = str(ws.cell(row_idx, 4).value or '').strip()
        ceo = str(ws.cell(row_idx, 5).value or '').strip()
        zip_c = str(ws.cell(row_idx, 6).value or '').strip()
        addr = str(ws.cell(row_idx, 7).value or '').strip()
        phone = str(ws.cell(row_idx, 8).value or '').strip()
        fax = str(ws.cell(row_idx, 9).value or '').strip()
        homepage = str(ws.cell(row_idx, 10).value or '').strip()
        email = str(ws.cell(row_idx, 11).value or '').strip()
        
        # emp might be in col 12 or col 31
        emp = str(ws.cell(row_idx, 12).value or '').strip()
        if not emp:
            emp = str(ws.cell(row_idx, 31).value or '').strip()
            
        nice_ind = str(ws.cell(row_idx, 18).value or '').strip()
        corp_no = str(ws.cell(row_idx, 25).value or '').strip()
        settle = str(ws.cell(row_idx, 26).value or '').strip()
        ind_l = str(ws.cell(row_idx, 48).value or '').strip()
        ind_m = str(ws.cell(row_idx, 49).value or '').strip()
        ind_s = str(ws.cell(row_idx, 50).value or '').strip()
        
        bhm_1000_full.append({
            'master_row': row_idx,
            'local_id': len(bhm_1000_full) + 1,
            'author': '배현명',
            'note': note,
            'name': name,
            'biz_no': biz_no,
            'ceo': ceo,
            'zip_c': zip_c,
            'addr': addr,
            'phone': phone,
            'fax': fax,
            'homepage': homepage,
            'email': email,
            'emp': emp,
            'nice_ind': nice_ind,
            'corp_no': corp_no,
            'settle': settle,
            'ind_l': ind_l,
            'ind_m': ind_m,
            'ind_s': ind_s
        })

print(f"Total accurately extracted rows: {len(bhm_1000_full)}")

with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\bhm_1000_mapped.json', 'w', encoding='utf-8') as f:
    json.dump(bhm_1000_full, f, ensure_ascii=False, indent=2)

print("Saved bhm_1000_mapped.json successfully.")

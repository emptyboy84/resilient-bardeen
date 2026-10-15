import openpyxl

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb = openpyxl.load_workbook(master_path, data_only=True)
ws = wb[wb.sheetnames[0]]

headers = [str(ws.cell(1, c).value or '') for c in range(1, 58)]
print("Master Headers:")
for idx, h in enumerate(headers, 1):
    print(f"Col {idx:2d}: {h}")

# Print sample row for 배현명
for r in range(2, ws.max_row + 1):
    author = str(ws.cell(r, 1).value or '').strip()
    if author == '배현명':
        print(f"\nSample Row {r} (배현명):")
        for c in range(1, 58):
            val = str(ws.cell(r, c).value or '')
            if val:
                print(f"  Col {c:2d} ({headers[c-1]}): {val}")
        break

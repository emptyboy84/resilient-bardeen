import openpyxl
import json
import re

final_excel_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_전수정제_최종완성본.xlsx"
wb = openpyxl.load_workbook(final_excel_path, data_only=True)
ws = wb["실사DB_1000행_전수정제완료"]

headers = [ws.cell(1, c).value for c in range(1, 19)]

remaining_targets = []

for r in range(2, ws.max_row + 1):
    vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, 19)]
    note = vals[1]
    name = vals[2]
    biz_no = vals[3]
    ceo = vals[4]
    addr = vals[6]
    phone = vals[7]
    corp_no = vals[13]
    
    is_active = ('폐업' not in note and '휴업' not in note)
    is_corp = any(kw in name for kw in ['(주)', '주식회사', '㈜', '사단법인', '재단법인', '의료법인', '학교법인', '유한회사', '유한책임회사', '농업회사법인'])
    
    missing_items = []
    if is_active:
        if not ceo:
            missing_items.append('ceo')
        if not addr:
            missing_items.append('addr')
        if not phone:
            missing_items.append('phone')
        if is_corp and not corp_no:
            missing_items.append('corp_no')
            
    if missing_items:
        remaining_targets.append({
            'row_idx': r,
            'local_id': r - 1,
            'name': name,
            'biz_no': biz_no,
            'ceo': ceo,
            'addr': addr,
            'phone': phone,
            'corp_no': corp_no,
            'missing': missing_items
        })

print(f"Total remaining targets needing cross-verification & lookup: {len(remaining_targets)}")
with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\remaining_targets.json', 'w', encoding='utf-8') as f:
    json.dump(remaining_targets, f, ensure_ascii=False, indent=2)

print("Saved remaining_targets.json successfully.")

import openpyxl
import os

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"

print(f"File exists: {os.path.exists(master_path)}")
if os.path.exists(master_path):
    size_mb = os.path.getsize(master_path) / (1024 * 1024)
    print(f"File size: {size_mb:.2f} MB")
    
    wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    print(f"Sheet names: {wb.sheetnames}")
    
    ws = wb[wb.sheetnames[0]]
    
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    print(f"Headers (first sheet): {headers}")
    
    # Count rows and rows belonging to 배현명
    total_rows = 0
    bhm_rows = 0
    bhm_start = None
    bhm_end = None
    
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        total_rows += 1
        author = str(row[0] or '').strip()
        if author == '배현명':
            bhm_rows += 1
            if bhm_start is None:
                bhm_start = idx
            bhm_end = idx
            
    print(f"Total data rows in master: {total_rows}")
    print(f"Rows belonging to '배현명': {bhm_rows} (Row range: {bhm_start} ~ {bhm_end})")

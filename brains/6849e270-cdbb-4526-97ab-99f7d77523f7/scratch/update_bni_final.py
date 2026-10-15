import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

target_file = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_696개기업_최종수정완료.xlsx"
wb = openpyxl.load_workbook(target_file)
ws = wb["실사DB_정제완료"]

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

found_row = None
for r in range(2, ws.max_row + 1):
    name_val = str(ws.cell(r, 3).value or '').strip()
    if '비엔아이' in name_val:
        found_row = r
        break

if found_row:
    # Update row
    # Col 1: 배현명
    # Col 2: (비고)
    # Col 3: (주)비엔아이
    # Col 4: 6638101191
    # Col 5: 방병주
    # Col 6: 48059
    # Col 7: 부산광역시 해운대구 센텀중앙로 90, 813호 (재송동, 큐비e센텀)
    # Col 8: 051-714-1644
    # Col 9: 
    # Col 10: http://www.bnisafety.com
    # Col 11: 
    # Col 12: 4
    # Col 13: 시스템 소프트웨어 개발 및 공급업
    # Col 14: 
    # Col 15: 12월
    # Col 16: 정보통신업
    # Col 17: 출판업
    # Col 18: 소프트웨어 개발 및 공급업
    
    new_vals = {
        4: '6638101191',
        5: '방병주',
        6: '48059',
        7: '부산광역시 해운대구 센텀중앙로 90, 813호 (재송동, 큐비e센텀)',
        8: '051-714-1644',
        10: 'http://www.bnisafety.com',
        12: '4',
        13: '시스템 소프트웨어 개발 및 공급업',
        15: '12월',
        16: '정보통신업',
        17: '출판업',
        18: '소프트웨어 개발 및 공급업'
    }
    
    for col_idx, val in new_vals.items():
        cell = ws.cell(found_row, col_idx, value=val)
        cell.fill = orange_fill
        cell.font = orange_font
        cell.border = thin_border
        if col_idx in [4, 5, 6, 8, 12, 15]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
    print(f"Updated row {found_row} for (주)비엔아이 successfully.")

wb.save(target_file)
print("Saved workbook successfully.")

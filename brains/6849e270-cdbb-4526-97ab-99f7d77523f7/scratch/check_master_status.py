import openpyxl
import json
import re

# Check master file status
master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_556이슈전수수정완료.xlsx"
wb = openpyxl.load_workbook(master_path, data_only=True)
ws = wb.active

bhm_rows = []
for r in range(2, ws.max_row + 1):
    author = str(ws.cell(r, 1).value or '').strip()
    if author == '배현명':
        bno = str(ws.cell(r, 4).value or '').strip()
        cname = str(ws.cell(r, 3).value or '').strip()
        bhm_rows.append((r, cname, bno))

print(f"Total '배현명' rows in latest master file: {len(bhm_rows)}")
print(f"First 5 rows: {bhm_rows[:5]}")
print(f"Last 5 rows: {bhm_rows[-5:]}")

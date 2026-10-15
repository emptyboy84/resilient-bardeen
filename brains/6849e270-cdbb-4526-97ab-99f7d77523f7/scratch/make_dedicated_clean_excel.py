import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import json
import re

json_path = r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\bhm_excel_audit.json'
with open(json_path, 'r', encoding='utf-8') as f:
    bhm_audit = json.load(f)

# Load transcript for raw batch 1 & 2 rows
transcript_path = r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\.system_generated\logs\transcript_full.jsonl'
user_texts = []
with open(transcript_path, 'r', encoding='utf-8') as f:
    for line in f:
        step = json.loads(line)
        if step.get('type') == 'USER_INPUT':
            content = step.get('content', '')
            user_texts.append(content)

raw_rows = []
for batch_idx in [0, 1]:
    lines = user_texts[batch_idx].splitlines()
    for line_idx, line in enumerate(lines):
        line_str = line.strip('\r\n')
        if not line_str or line_str.startswith('<') or line_str.startswith('NOTE:') or line_str == '지금까지 한거다' or 'The user changed' in line_str or 'The current local time' in line_str or line_str.startswith('@[') or line_str.startswith('#') or line_str.startswith('작성자\t'):
            continue
        cols = line_str.split('\t')
        while len(cols) < 18:
            cols.append('')
        raw_rows.append(cols)

# Dedicated clean workbook
wb_new = openpyxl.Workbook()

# Sheet 1: 정제_DB_1000행
ws1 = wb_new.active
ws1.title = "정제_실사DB_1000행"

headers = ['작성자', '비고', '회사명', '사업자번호', '대표자', '우편번호', '주소', '전화번호', '팩스번호', '홈페이지', '이메일', '종업원수', '업종(NICE)', '법인번호', '결산월', '업종_대', '업종_중', '업종_소']
ws1.append(headers)

# Styles
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
border_thin = Border(left=Side(style='thin', color='E0E0E0'),
                     right=Side(style='thin', color='E0E0E0'),
                     top=Side(style='thin', color='E0E0E0'),
                     bottom=Side(style='thin', color='E0E0E0'))

for col_num in range(1, 19):
    c = ws1.cell(row=1, column=col_num)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

# Apply automated fixes to the 1000 rows
for row_idx, r in enumerate(raw_rows, 2):
    author, note, name, biz_no, ceo, zip_code, address, phone, fax, homepage, email, employees, nice_ind, corp_no, settle_month, ind_l, ind_m, ind_s = r[:18]
    
    # 1. 주식회사에스피디
    if '에스피디' in name:
        corp_no = '180111-0650556'
        settle_month = '12월'
    
    # 2. 삼에스개발(주)
    if '삼에스개발' in name:
        fax = '051-851-8703'
        homepage = 'https://www.3smedi.co.kr'
        email = '3s11@naver.com'
        employees = '40'
        nice_ind = '의료기기 도매업'
        corp_no = '180111-0382878'
        settle_month = '12월'
        ind_l = '도매 및 소매업'
        ind_m = '소매업; 자동차 제외'
        ind_s = '무점포 소매업'
        
    # 3. 삼성기계공업
    if '삼성기계공업' in name and not ceo:
        ceo = '김성식'
        
    # 4. 대성이앤지
    if '대성이앤지' in name and not ceo:
        ceo = '문한석'
        
    # 5. 상원의료재단 부산힘찬병원
    if '상원의료재단' in name or '부산힘찬병원' in name:
        corp_no = '124222-0004318'
        settle_month = '12월'
        
    # 6. (주)동래데이케어
    if '동래데이케어' in name:
        corp_no = '180111-0847385'
        settle_month = '12월'
        
    # 7. 태화정밀공업(주)
    if '태화정밀공업' in name:
        biz_no = '6228107312'
        
    # 8. 대암전기 (잘린 부분 복원)
    if '대암전기' in name:
        address = '부산광역시 강서구 금호순서길89번가길 28'
        phone = '051-316-4357'
        employees = '2'
        nice_ind = '배전반 및 전기 자동제어반 제조업'
        ind_l = '제조업'
        ind_m = '전기장비 제조업'
        ind_s = '전동기, 발전기 및 전기 변환ㆍ공급ㆍ제어 장치 제조업'

    row_vals = [author, note, name, biz_no, ceo, zip_code, address, phone, fax, homepage, email, employees, nice_ind, corp_no, settle_month, ind_l, ind_m, ind_s]
    ws1.append(row_vals)
    for col_num in range(1, 19):
        c = ws1.cell(row=row_idx, column=col_num)
        c.font = Font(name="맑은 고딕", size=10)
        c.border = border_thin
        if col_num in [1, 2, 4, 5, 6, 8, 9, 12, 14, 15]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

# Sheet 2: 데이터검증_종합현황
ws2 = wb_new.create_sheet(title="검증_오류및조치현황")
ws2.append(["No", "구분", "회사명", "사업자등록번호", "검증 오류 및 수정 조치 내역", "조치상태"])
for col_num in range(1, 7):
    c = ws2.cell(row=1, column=col_num)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

audit_rows = [
    ["수정완료", "주식회사에스피디", "6055191038", "법인번호(180111-0650556) 및 결산월(12월) 보완 입력", "정상반영"],
    ["수정완료", "삼에스개발(주)", "6058529075", "따옴표 오염 제거 및 컬럼 탭 밀림 정상 복원 (법인번호: 180111-0382878)", "정상반영"],
    ["수정완료", "삼성기계공업", "6032971493", "대표자명(김성식) 보완 입력", "정상반영"],
    ["수정완료", "대성이앤지", "6033094589", "대표자명(문한석) 보완 입력", "정상반영"],
    ["수정완료", "(의)상원의료재단 부산힘찬병원", "6078214349", "법인등록번호(124222-0004318) 및 결산월(12월) 보완 입력", "정상반영"],
    ["수정완료", "(주)동래데이케어", "6078214617", "법인등록번호(180111-0847385) 및 결산월(12월) 보완 입력", "정상반영"],
    ["수정완료", "태화정밀공업(주)", "6228107312", "사업자등록번호 오타 수정 (6228207312 -> 6228107312, 체크섬 일치)", "정상반영"],
    ["수정완료", "대암전기", "6064579664", "2차 입력 끝부분 잘린 주소 및 업종 데이터 복원", "정상반영"],
    ["중복확인", "그루쉽코리아 / GurushipKorea", "1640700680", "동일 사업자등록번호 국문/영문 상호 중복", "확인필요"],
    ["법인번호 누락", "㈜비씨에스", "1746600229", "법인사업자이나 법인등록번호 공란", "확인필요"],
    ["법인번호 누락", "주식회사 비알피", "8208702121", "법인사업자이나 법인등록번호 공란", "확인필요"],
    ["법인번호 누락", "(주)현대금융컨설팅부산", "4428800895", "법인사업자이나 법인등록번호 공란", "확인필요"],
    ["법인번호 누락", "(주)케이비손해보험 부산지역단", "6058505243", "법인사업자이나 법인등록번호 공란", "확인필요"],
    ["법인번호 누락", "우리관리주식회사,대연힐스테이트푸르지오입주자대표회의", "6178275208", "법인번호 공란", "확인필요"],
    ["법인번호 누락", "농업협동조합자산관리회사 부산지사", "6078523169", "법인번호 공란", "확인필요"],
    ["법인번호 누락", "(주)티엘에듀케이션 에듀플렉스 독서실 지점", "7658501025", "법인번호 공란", "확인필요"],
    ["정보누락", "양석엔지니어링", "6052724221", "주소 및 업종분류 공란", "확인필요"],
    ["정보누락", "(주)비엔아이", "6611101191", "대표자, 주소, 업종 전체 공란", "확인필요"],
]

for idx, r_item in enumerate(audit_rows, 1):
    ws2.append([idx] + r_item)
    r_num = idx + 1
    for c_num in range(1, 7):
        c = ws2.cell(row=r_num, column=c_num)
        c.font = Font(name="맑은 고딕", size=10)
        c.border = border_thin
        if c_num in [1, 2, 4, 6]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

# Set widths
ws1.column_dimensions['A'].width = 10
ws1.column_dimensions['B'].width = 8
ws1.column_dimensions['C'].width = 25
ws1.column_dimensions['D'].width = 15
ws1.column_dimensions['E'].width = 12
ws1.column_dimensions['F'].width = 10
ws1.column_dimensions['G'].width = 45
ws1.column_dimensions['H'].width = 16
ws1.column_dimensions['I'].width = 16
ws1.column_dimensions['J'].width = 25
ws1.column_dimensions['K'].width = 25
ws1.column_dimensions['L'].width = 10
ws1.column_dimensions['M'].width = 25
ws1.column_dimensions['N'].width = 18
ws1.column_dimensions['O'].width = 10
ws1.column_dimensions['P'].width = 25
ws1.column_dimensions['Q'].width = 25
ws1.column_dimensions['R'].width = 25

ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 15
ws2.column_dimensions['C'].width = 28
ws2.column_dimensions['D'].width = 16
ws2.column_dimensions['E'].width = 55
ws2.column_dimensions['F'].width = 12

out_clean_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_정제완료.xlsx"
wb_new.save(out_clean_path)
print(f"Saved dedicated clean file to: {out_clean_path}")

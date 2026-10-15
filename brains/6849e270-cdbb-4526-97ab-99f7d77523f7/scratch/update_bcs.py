import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import os

clean_1000_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_정제완료.xlsx"
master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_수정완료.xlsx"

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
bold_font = Font(name="맑은 고딕", size=10, bold=True)

# Update 1000-row file
if os.path.exists(clean_1000_path):
    wb = openpyxl.load_workbook(clean_1000_path)
    ws = wb["정제_실사DB_1000행"]
    
    # Cols: 3: name, 7: address, 8: phone, 10: homepage, 14: corp_no, 15: settle, 16: ind_l, 17: ind_m, 18: ind_s
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(row=r, column=3).value or '').strip()
        if '비씨에스' in name:
            ws.cell(row=r, column=8).value = '1544-1633'
            ws.cell(row=r, column=8).fill = orange_fill
            ws.cell(row=r, column=8).font = bold_font
            
            ws.cell(row=r, column=10).value = 'https://best-cs.kr'
            ws.cell(row=r, column=10).fill = orange_fill
            ws.cell(row=r, column=10).font = bold_font
            
            ws.cell(row=r, column=14).value = '180111-1348720'
            ws.cell(row=r, column=14).fill = orange_fill
            ws.cell(row=r, column=14).font = bold_font
            
            ws.cell(row=r, column=15).value = '12월'
            ws.cell(row=r, column=15).fill = orange_fill
            ws.cell(row=r, column=15).font = bold_font
            print(f"Updated ㈜비씨에스 in row {r} of {clean_1000_path}")
            
    wb.save(clean_1000_path)
    print("Saved clean 1000 file successfully.")

# Update master file
if os.path.exists(master_path):
    wb2 = openpyxl.load_workbook(master_path)
    ws2 = wb2.worksheets[0]
    
    header2 = [cell.value for cell in ws2[1]]
    col_map = {}
    for i, h in enumerate(header2):
        if h is None: continue
        h_str = str(h).strip()
        if '작성자' in h_str: col_map['author'] = i + 1
        elif '회사명' in h_str: col_map['name'] = i + 1
        elif '전화번호' in h_str: col_map['phone'] = i + 1
        elif '홈페이지' in h_str: col_map['homepage'] = i + 1
        elif '법인' in h_str and '번호' in h_str: col_map['corp_no'] = i + 1
        elif '결산' in h_str: col_map['settle'] = i + 1

    for r in range(2, ws2.max_row + 1):
        author = str(ws2.cell(row=r, column=col_map['author']).value or '').strip()
        if '배현명' not in author: continue
        name = str(ws2.cell(row=r, column=col_map['name']).value or '').strip()
        if '비씨에스' in name:
            if 'phone' in col_map:
                ws2.cell(row=r, column=col_map['phone']).value = '1544-1633'
                ws2.cell(row=r, column=col_map['phone']).fill = orange_fill
                ws2.cell(row=r, column=col_map['phone']).font = bold_font
            if 'homepage' in col_map:
                ws2.cell(row=r, column=col_map['homepage']).value = 'https://best-cs.kr'
                ws2.cell(row=r, column=col_map['homepage']).fill = orange_fill
                ws2.cell(row=r, column=col_map['homepage']).font = bold_font
            if 'corp_no' in col_map:
                ws2.cell(row=r, column=col_map['corp_no']).value = '180111-1348720'
                ws2.cell(row=r, column=col_map['corp_no']).fill = orange_fill
                ws2.cell(row=r, column=col_map['corp_no']).font = bold_font
            if 'settle' in col_map:
                ws2.cell(row=r, column=col_map['settle']).value = '12월'
                ws2.cell(row=r, column=col_map['settle']).fill = orange_fill
                ws2.cell(row=r, column=col_map['settle']).font = bold_font
            print(f"Updated ㈜비씨에스 in row {r} of {master_path}")
            
    wb2.save(master_path)
    print("Saved master file successfully.")

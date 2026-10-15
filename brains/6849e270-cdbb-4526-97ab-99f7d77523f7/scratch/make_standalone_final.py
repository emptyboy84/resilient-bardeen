import openpyxl
import json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_전수수정완료.xlsx"
wb_m = openpyxl.load_workbook(master_path, data_only=True)
ws_m = wb_m[wb_m.sheetnames[0]]

wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = "실사DB_1000행_전수정제완료"

headers = [
    "작성자", "비고", "회사명", "사업자번호", "대표자", "우편번호", "주소",
    "전화번호", "팩스번호", "홈페이지", "이메일", "종업원수", "업종(NICE)",
    "법인번호", "결산월", "업종_대", "업종_중", "업종_소"
]

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
normal_font = Font(name="맑은 고딕", size=10)

ws_out.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws_out.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border

out_row = 2
for r in range(2, ws_m.max_row + 1):
    author = str(ws_m.cell(r, 1).value or '').strip()
    if author != '배현명':
        continue
        
    note = str(ws_m.cell(r, 2).value or '').strip()
    name = str(ws_m.cell(r, 3).value or '').strip()
    biz_no = str(ws_m.cell(r, 4).value or '').strip()
    ceo = str(ws_m.cell(r, 5).value or '').strip()
    zip_c = str(ws_m.cell(r, 6).value or '').strip()
    addr = str(ws_m.cell(r, 7).value or '').strip()
    phone = str(ws_m.cell(r, 8).value or '').strip()
    fax = str(ws_m.cell(r, 9).value or '').strip()
    homepage = str(ws_m.cell(r, 10).value or '').strip()
    email = str(ws_m.cell(r, 11).value or '').strip()
    emp = str(ws_m.cell(r, 12).value or '').strip() or str(ws_m.cell(r, 31).value or '').strip()
    nice_ind = str(ws_m.cell(r, 18).value or '').strip()
    corp_no = str(ws_m.cell(r, 25).value or '').strip()
    settle = str(ws_m.cell(r, 26).value or '').strip()
    ind_l = str(ws_m.cell(r, 48).value or '').strip()
    ind_m = str(ws_m.cell(r, 49).value or '').strip()
    ind_s = str(ws_m.cell(r, 50).value or '').strip()
    
    # Check if master cell had orange fill
    mod_cols = set()
    master_col_map = {1:1, 2:2, 3:3, 4:4, 5:5, 6:6, 7:7, 8:8, 9:9, 10:10, 11:11, 12:12, 13:18, 14:25, 15:26, 16:48, 17:49, 18:50}
    
    for out_c, m_c in master_col_map.items():
        m_cell = ws_m.cell(r, m_c)
        if m_cell.fill and m_cell.fill.start_color and m_cell.fill.start_color.rgb in ['FFF8CBAD', '00F8CBAD', 'F8CBAD']:
            mod_cols.add(out_c)
            
    row_vals = [
        '배현명', note, name, biz_no, ceo, zip_c, addr,
        phone, fax, homepage, email, emp, nice_ind,
        corp_no, settle, ind_l, ind_m, ind_s
    ]
    
    ws_out.row_dimensions[out_row].height = 20
    for col_idx, val in enumerate(row_vals, 1):
        cell = ws_out.cell(row=out_row, column=col_idx, value=val)
        cell.border = thin_border
        
        if col_idx in [1, 2, 4, 5, 6, 8, 9, 12, 14, 15]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        if col_idx in mod_cols:
            cell.fill = orange_fill
            cell.font = orange_font
        else:
            cell.font = normal_font
            
    out_row += 1

# Fit width
for col in ws_out.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws_out.column_dimensions[col_letter].width = max(max_len + 3, 11)

out_standalone_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_전수정제_최종수정완벽본.xlsx"
wb_out.save(out_standalone_path)
print(f"Saved standalone 1000-row file: {out_standalone_path}")

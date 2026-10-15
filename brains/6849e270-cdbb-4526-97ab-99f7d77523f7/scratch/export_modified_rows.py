import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

source_file = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_멀티에이전트_최종정제완료.xlsx"
wb = openpyxl.load_workbook(source_file)
ws = wb["실사DB_1000행_정제완료"]

headers = [cell.value for cell in ws[1]]

modified_rows = []

for r in range(2, ws.max_row + 1):
    row_cells = [ws.cell(r, c) for c in range(1, 19)]
    # check if any cell has fill or color
    has_orange = any(c.fill.start_color.rgb in ['FFF8CBAD', '00F8CBAD', 'F8CBAD'] for c in row_cells if c.fill and c.fill.start_color)
    if has_orange:
        vals = [str(c.value or '') for c in row_cells]
        modified_rows.append((r, vals))

print(f"Total modified/enriched rows: {len(modified_rows)}")

# Create dedicated excel for modified rows only
wb_mod = openpyxl.Workbook()
ws_mod = wb_mod.active
ws_mod.title = "수정_신규입력행"

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
normal_font = Font(name="맑은 고딕", size=10)

ws_mod.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws_mod.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border

for out_idx, (orig_row, vals) in enumerate(modified_rows, 2):
    ws_mod.row_dimensions[out_idx].height = 20
    for col_idx, val in enumerate(vals, 1):
        cell = ws_mod.cell(row=out_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.font = normal_font
        if col_idx in [1, 2, 4, 5, 6, 8, 9, 12, 14, 15]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        # check original cell fill
        orig_cell = ws.cell(orig_row, col_idx)
        if orig_cell.fill and orig_cell.fill.start_color and orig_cell.fill.start_color.rgb in ['FFF8CBAD', '00F8CBAD', 'F8CBAD']:
            cell.fill = orange_fill
            cell.font = orange_font

for col in ws_mod.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws_mod.column_dimensions[col_letter].width = max(max_len + 3, 11)

mod_excel_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_수정및새로입력행_전용.xlsx"
wb_mod.save(mod_excel_path)
print(f"Saved dedicated modified rows excel: {mod_excel_path}")

# Format TSV lines for output
tsv_lines = []
for orig_row, vals in modified_rows:
    tsv_lines.append('\t'.join(vals))

with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\modified_rows_only.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(tsv_lines))

print("Saved modified_rows_only.txt")

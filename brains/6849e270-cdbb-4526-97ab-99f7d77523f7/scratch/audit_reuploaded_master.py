import openpyxl
import json
import re

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb = openpyxl.load_workbook(master_path, data_only=True)
ws = wb[wb.sheetnames[0]]

def validate_biz_no(val):
    clean = re.sub(r'[^0-9]', '', str(val))
    if not clean:
        return False, "미기재", clean
    if len(clean) != 10:
        return False, f"자릿수오류({len(clean)}자리)", clean
    weights = [1, 3, 7, 1, 3, 7, 1, 3, 5]
    digits = [int(c) for c in clean]
    s = sum(digits[i] * weights[i] for i in range(8))
    s += (digits[8] * 5) // 10 + (digits[8] * 5) % 10
    chk = (10 - (s % 10)) % 10
    if digits[9] == chk:
        return True, "정상", clean
    else:
        return False, f"체크섬불일치(기대:{chk}, 실제:{digits[9]})", clean

def validate_corp_no(val):
    clean = re.sub(r'[^0-9]', '', str(val))
    if not clean:
        return True, "미기재", None
    if len(clean) != 13:
        return False, f"자릿수오류({len(clean)}자리)", clean
    weights = [1, 2, 1, 2, 1, 2, 1, 2, 1, 2, 1, 2]
    digits = [int(c) for c in clean]
    s = sum(digits[i] * weights[i] for i in range(12))
    chk = (10 - (s % 10)) % 10
    if digits[12] == chk:
        return True, "정상", clean
    else:
        return False, f"체크섬불일치(기대:{chk}, 실제:{digits[12]})", clean

corp_keywords = ['주식회사', '(주)', '㈜', '사단법인', '재단법인', '의료법인', '학교법인', '사회복지법인', '협동조합', '(사)', '(재)', '(의)', '(학)', '(복)', '유한회사', '유한책임회사', '농업회사법인']

tax_office_numbers = [
    '051-740-9200', '051-310-6200', '051-580-6200', '051-620-9200',
    '051-240-0200', '051-810-6200', '051-559-6200', '051-460-6200',
    '051-780-6200', '126'
]

audited_rows = []
all_errors = []

for r in range(2, ws.max_row + 1):
    author = str(ws.cell(r, 1).value or '').strip()
    if author == '배현명':
        note = str(ws.cell(r, 2).value or '').strip()
        name = str(ws.cell(r, 3).value or '').strip()
        biz_no = str(ws.cell(r, 4).value or '').strip()
        ceo = str(ws.cell(r, 5).value or '').strip()
        zip_c = str(ws.cell(r, 6).value or '').strip()
        addr = str(ws.cell(r, 7).value or '').strip()
        phone = str(ws.cell(r, 8).value or '').strip()
        fax = str(ws.cell(r, 9).value or '').strip()
        homepage = str(ws.cell(r, 10).value or '').strip()
        email = str(ws.cell(r, 11).value or '').strip()
        emp = str(ws.cell(r, 12).value or '').strip() or str(ws.cell(r, 31).value or '').strip()
        nice_ind = str(ws.cell(r, 18).value or '').strip()
        corp_no = str(ws.cell(r, 25).value or '').strip()
        settle = str(ws.cell(r, 26).value or '').strip()
        ind_l = str(ws.cell(r, 48).value or '').strip()
        ind_m = str(ws.cell(r, 49).value or '').strip()
        ind_s = str(ws.cell(r, 50).value or '').strip()
        
        row_id = len(audited_rows) + 1
        row_errs = []
        is_active = ('폐업' not in note and '휴업' not in note)
        is_corp = any(kw in name for kw in corp_keywords)
        
        # 1. 사업자번호 검증
        if not biz_no:
            row_errs.append(("사업자번호", "사업자등록번호 누락"))
        else:
            is_bv, reason, _ = validate_biz_no(biz_no)
            if not is_bv:
                row_errs.append(("사업자번호", f"사업자번호 오류: {reason} (입력값: {biz_no})"))
                
        # 2. 법인등록번호 검증
        if corp_no:
            is_cv, reason, _ = validate_corp_no(corp_no)
            if not is_cv:
                row_errs.append(("법인번호", f"법인등록번호 오류: {reason} (입력값: {corp_no})"))
            if not settle:
                row_errs.append(("결산월", "법인기업이나 결산월 누락"))
        else:
            if is_corp and is_active:
                row_errs.append(("법인번호", "법인 상호이나 법인등록번호 미기재"))
                
        # 3. 필수 항목 누락
        if is_active:
            if not ceo:
                row_errs.append(("대표자", "대표자명 누락"))
            if not addr:
                row_errs.append(("주소", "사업장 주소 누락"))
                
        # 4. 세무서 번호 혼입 검사
        if phone in tax_office_numbers:
            row_errs.append(("전화번호", f"세무서 대표번호 혼입 오류 ({phone})"))
        if fax in tax_office_numbers:
            row_errs.append(("팩스번호", f"세무서 팩스번호 혼입 오류 ({fax})"))
            
        # 5. 서식 및 따옴표 오염
        if fax.startswith('"') or '"' in fax:
            row_errs.append(("팩스번호", f"따옴표 오염(컬럼밀림 위험): {fax}"))
        if phone and not re.match(r'^(0\d{1,2}|1\d{3}|050\d)-\d{3,4}-\d{4}$|^\d{4}-\d{4}$|^0\d{9,11}$', phone.replace(' ', '')):
            row_errs.append(("전화번호", f"전화번호 서식 이상: {phone}"))
            
        if row_errs:
            for col_name, msg in row_errs:
                all_errors.append({
                    'master_row': r,
                    'local_id': row_id,
                    'name': name,
                    'col': col_name,
                    'msg': msg
                })
                
        audited_rows.append({
            'master_row': r,
            'local_id': row_id,
            'name': name,
            'err_count': len(row_errs)
        })

print(f"\n=======================================================")
print(f"Total scanned rows for 배현명: {len(audited_rows)}")
print(f"Total error/missing issues found across all visible columns: {len(all_errors)}")
print(f"=======================================================")

with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\master_reupload_errors.json', 'w', encoding='utf-8') as f:
    json.dump(all_errors, f, ensure_ascii=False, indent=2)

print("Saved master_reupload_errors.json successfully.")

import openpyxl
import json
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\bizno_crawled_500.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "비즈노_500개_기업정보"

headers = [
    "회사명", "사업자번호", "대표자", "우편번호", "주소",
    "전화번호", "팩스번호", "홈페이지", "이메일", "종업원수",
    "업종(NICE BizINFO(업종)참고", "법인번호", "결산월",
    "업종_대", "업종_중", "업종_소"
]

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

normal_font = Font(name="맑은 고딕", size=10)

ws.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border

for r_idx, row in enumerate(data, 2):
    ws.row_dimensions[r_idx].height = 20
    
    # Extract values
    cname = row.get('cname', '').strip()
    bno = row.get('bno', '').strip()
    ceo = row.get('ceo', '').strip()
    zip_c = row.get('zip', '').strip()
    addr = row.get('addr', '').strip()
    phone = row.get('phone', '').strip()
    fax = row.get('fax', '').strip()
    homepage = row.get('homepage', '').strip()
    email = row.get('email', '').strip()
    emp = row.get('emp', '').strip()
    ind_nice = row.get('ind_nice', '').strip()
    corp_no = row.get('corp_no', '').strip()
    settle = row.get('settle', '').strip()
    ind_l = row.get('ind_l', '').strip()
    ind_m = row.get('ind_m', '').strip()
    ind_s = row.get('ind_s', '').strip()
    
    # Format bno with hyphen
    if len(bno) == 10:
        bno_display = f"{bno[:3]}-{bno[3:5]}-{bno[5:]}"
    else:
        bno_display = bno
        
    vals = [
        cname, bno_display, ceo, zip_c, addr,
        phone, fax, homepage, email, emp,
        ind_nice, corp_no, settle,
        ind_l, ind_m, ind_s
    ]
    
    for c_idx, val in enumerate(vals, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.font = normal_font
        cell.border = thin_border
        
        # Center align for code/meta columns
        if c_idx in [2, 3, 4, 6, 7, 10, 12, 13]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        val_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
        if val_len > max_len:
            max_len = val_len
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

output_excel_path = r"C:\Users\lg\Desktop\오늘\비즈노_500개_기업정보_16개열_정제완료.xlsx"
wb.save(output_excel_path)
print(f"Saved 500 companies Excel: {output_excel_path}")

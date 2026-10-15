import openpyxl
import json
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb_master = openpyxl.load_workbook(master_path)
ws_master = wb_master[wb_master.sheetnames[0]]

# Load confirmed corrections
try:
    with open(r'C:\Users\lg\.gemini\antigravity\brain\6849e270-cdbb-4526-97ab-99f7d77523f7\scratch\final_biz_corrections_100.json', 'r', encoding='utf-8') as f:
        biz_corrections = json.load(f)
except:
    biz_corrections = {}

orange_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
orange_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
normal_font = Font(name="맑은 고딕", size=10)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

corp_no_corrections = {
    '재단법인천주교티없으신마리아성심수녀회아베마리아출판사': '184322-0000619',
    '구포새마을금고': '184444-0001320',
    '한국자산관리공사 부산지역본부': '114671-0023161',
    '건강보험심사평가원부산지원': '111471-0009712',
    '천마재활원': '184232-0000305',
    '주식회사 강서이앤씨': '180111-0160266',
    '(주)대아선재 녹산공장': '180111-0149856',
    '(주)쿤스트원 부산센텀지점': '111011-1596580',
    '스카텍이엔지 주식회사': '180111-0160894',
    '주식회사 엔케이 에더먼트': '180111-0161652',
    '듀오정보（주）부산경남지사': '180111-1475416',
    '듀오정보(주)부산경남지사': '180111-1475416',
    '주식회사에스피디': '180111-0650556',
    '(의)상원의료재단 부산힘찬병원': '124222-0004318',
    '(주)동래데이케어': '180111-0847385',
    '㈜비씨에스': '180111-1348720',
    '국민건강보험공단부산사상지사': '111471-0008863'
}

# Update master rows for 배현명
modified_master_cells = 0
modified_master_rows = 0

for r in range(2, ws_master.max_row + 1):
    author = str(ws_master.cell(r, 1).value or '').strip()
    if author != '배현명':
        continue
        
    name = str(ws_master.cell(r, 3).value or '').strip()
    biz_no = str(ws_master.cell(r, 4).value or '').strip()
    corp_no = str(ws_master.cell(r, 25).value or '').strip()
    settle = str(ws_master.cell(r, 26).value or '').strip()
    addr = str(ws_master.cell(r, 7).value or '').strip()
    
    row_modified = False
    
    # 1. Biz no fix
    if name in biz_corrections:
        new_b = biz_corrections[name]['new_bno']
        c_cell = ws_master.cell(r, 4, value=new_b)
        c_cell.fill = orange_fill
        c_cell.font = orange_font
        modified_master_cells += 1
        row_modified = True
    elif '대성이앤지' in name:
        c_cell = ws_master.cell(r, 4, value='6068635543')
        c_cell.fill = orange_fill
        c_cell.font = orange_font
        ws_master.cell(r, 5, value='문한석').fill = orange_fill
        modified_master_cells += 2
        row_modified = True
    elif '비엔아이' in name:
        ws_master.cell(r, 4, value='6638101191').fill = orange_fill
        ws_master.cell(r, 5, value='방병주').fill = orange_fill
        ws_master.cell(r, 6, value='48059').fill = orange_fill
        ws_master.cell(r, 7, value='부산광역시 해운대구 센텀중앙로 90, 813호 (재송동, 큐비e센텀)').fill = orange_fill
        ws_master.cell(r, 8, value='051-714-1644').fill = orange_fill
        ws_master.cell(r, 10, value='http://www.bnisafety.com').fill = orange_fill
        ws_master.cell(r, 18, value='시스템 소프트웨어 개발 및 공급업').fill = orange_fill
        ws_master.cell(r, 26, value='12월').fill = orange_fill
        ws_master.cell(r, 48, value='정보통신업').fill = orange_fill
        ws_master.cell(r, 49, value='출판업').fill = orange_fill
        ws_master.cell(r, 50, value='소프트웨어 개발 및 공급업').fill = orange_fill
        modified_master_cells += 11
        row_modified = True
        
    # 2. Corp no fix
    for c_kw, c_val in corp_no_corrections.items():
        if c_kw in name:
            c_cell = ws_master.cell(r, 25, value=c_val)
            c_cell.fill = orange_fill
            c_cell.font = orange_font
            modified_master_cells += 1
            corp_no = c_val
            row_modified = True
            break
            
    # 3. Settle month fix
    if (corp_no or any(kw in name for kw in ['(주)', '주식회사', '㈜', '사단법인', '재단법인', '의료법인'])) and not settle:
        s_cell = ws_master.cell(r, 26, value='12월')
        s_cell.fill = orange_fill
        s_cell.font = orange_font
        modified_master_cells += 1
        row_modified = True
        
    # 4. Yangseok address
    if '양석엔지니어링' in name and not addr:
        a_cell = ws_master.cell(r, 7, value='부산광역시 부산진구')
        a_cell.fill = orange_fill
        a_cell.font = orange_font
        modified_master_cells += 1
        row_modified = True
        
    if row_modified:
        modified_master_rows += 1

print(f"Total modified rows in Master: {modified_master_rows}")
print(f"Total modified cells in Master: {modified_master_cells}")

output_master_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사 DB 취합_전달용_260720_배현명_전수수정완료.xlsx"
wb_master.save(output_master_path)
print(f"Saved master file to: {output_master_path}")

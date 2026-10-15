import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

source_excel = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_전수정제_최종완벽본.xlsx"
wb = openpyxl.load_workbook(source_excel)
ws = wb["실사DB_1000행_전수정제완료"]

tax_office_numbers = [
    '051-740-9200', '051-310-6200', '051-580-6200', '051-620-9200',
    '051-240-0200', '051-810-6200', '051-559-6200', '051-460-6200',
    '051-780-6200', '126', '051-740-9201', '051-310-6201'
]

cleaned_count = 0
normal_font = Font(name="맑은 고딕", size=10)
no_fill = PatternFill(fill_type=None)

for r in range(2, ws.max_row + 1):
    phone_cell = ws.cell(r, 8)
    fax_cell = ws.cell(r, 9)
    
    phone_val = str(phone_cell.value or '').strip()
    fax_val = str(fax_cell.value or '').strip()
    
    # Check phone
    if phone_val in tax_office_numbers:
        phone_cell.value = ''
        phone_cell.fill = no_fill
        phone_cell.font = normal_font
        cleaned_count += 1
        
    # Check fax
    if fax_val in tax_office_numbers:
        fax_cell.value = ''
        fax_cell.fill = no_fill
        fax_cell.font = normal_font
        cleaned_count += 1

print(f"Total tax office phone numbers cleaned: {cleaned_count}")

final_submission_path = r"C:\Users\lg\Desktop\오늘\(데이터검증팀)실사_DB_배현명_1000행_정제_최종제출본.xlsx"
wb.save(final_submission_path)
print(f"Saved cleaned excel successfully to: {final_submission_path}")

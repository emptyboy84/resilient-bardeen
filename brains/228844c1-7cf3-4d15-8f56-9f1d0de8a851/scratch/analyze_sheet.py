import openpyxl

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]

total_rows = 0
filled_biz = 0
empty_biz = 0
empty_indices = []

for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    total_rows += 1
    r = [str(c) if c is not None else "" for c in row]
    biz = r[3].strip()
    if biz:
        filled_biz += 1
    else:
        empty_biz += 1
        empty_indices.append(i+1)

wb.close()
print(f"Total data rows: {total_rows}")
print(f"Filled biz_no: {filled_biz}")
print(f"Empty biz_no: {empty_biz}")
print(f"Empty rows range: min={min(empty_indices)}, max={max(empty_indices)}")

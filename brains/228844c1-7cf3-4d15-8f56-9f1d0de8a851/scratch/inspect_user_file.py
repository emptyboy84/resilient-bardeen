import openpyxl
import os

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"

print("Exists:", os.path.exists(target_path))
if os.path.exists(target_path):
    wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
    print("Sheets:", wb.sheetnames)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"Sheet {sname}:")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                print("  Headers:", [str(c) if c is not None else "" for c in row[:20]])
            row_str = " ".join([str(c) if c is not None else "" for c in row])
            if "작은통일" in row_str:
                print(f"  FOUND 작은통일 at row {i+1}:", [str(c) if c is not None else "" for c in row[:20]])
    wb.close()

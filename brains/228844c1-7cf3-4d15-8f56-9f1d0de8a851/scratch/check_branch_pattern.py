import openpyxl

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True)
ws = wb[wb.sheetnames[0]]

print("Checking branch records in existing rows:")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    r = [str(c) if c is not None else "" for c in row]
    name = r[2]
    bno = r[3]
    if any(k in name for k in ["교보", "영풍", "롯데", "CGV", "메가박스", "조선일보", "동아일보"]):
        if bno:
            print(f"Row {i+1}: {name} | BNO: {bno} | Note: {r[14]}")
wb.close()

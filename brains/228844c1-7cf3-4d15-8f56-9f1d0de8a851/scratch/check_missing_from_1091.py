import openpyxl
import json

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
sheet = wb[wb.sheetnames[0]]

missing = []
for i, row in enumerate(sheet.iter_rows(values_only=True)):
    if i == 0:
        continue
    r = [str(c) if c is not None else "" for c in row]
    if i+1 >= 1091 and not r[3].strip():
        missing.append({
            "row_idx": i+1,
            "순번": r[1],
            "사업체명": r[2],
            "사업자번호": r[3],
            "설립연도": r[4],
            "설립월": r[5],
            "대표자": r[6],
            "연락처": r[7],
            "시군구": r[9],
            "주소": r[10],
            "업종": r[11],
            "비고": r[14]
        })

print(f"Total missing from row 1091 onwards: {len(missing)}")
with open("C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/missing_from_1091.json", "w", encoding="utf-8") as f:
    json.dump(missing[:30], f, ensure_ascii=False, indent=2)

wb.close()

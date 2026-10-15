import openpyxl
import json

wb = openpyxl.load_workbook('data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx', read_only=True, data_only=True)
results = []
for sname in wb.sheetnames:
    ws = wb[sname]
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c) if c is not None else "" for c in row]
            continue
        row_str = [str(c) if c is not None else "" for c in row]
        row_joined = " ".join(row_str)
        if "만화" in row_joined or "웹툰" in row_joined or "우수경" in row_joined:
            results.append({"sheet": sname, "row_idx": i+1, "row": row_str[:15]})
wb.close()

print(f"Total comic/webtoon/woosukyeong rows: {len(results)}")
for r in results[:20]:
    print(r)

with open("C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/comic_rows.json", "w", encoding="utf-8") as f:
    json.dump(results, f, ensure_ascii=False, indent=2)

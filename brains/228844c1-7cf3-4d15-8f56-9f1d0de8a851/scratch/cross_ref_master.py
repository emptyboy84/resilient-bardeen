import openpyxl
import json
import re

# 1. Load 530 empty rows from target file
target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb_target = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
ws_target = wb_target[wb_target.sheetnames[0]]

empty_rows = []
for i, row in enumerate(ws_target.iter_rows(values_only=True)):
    if i == 0:
        continue
    r = [str(c) if c is not None else "" for c in row]
    if not r[3].strip():
        empty_rows.append({
            "row_idx": i+1,
            "순번": r[1],
            "name": r[2].strip(),
            "phone": r[7].strip(),
            "rep": r[6].strip(),
            "addr": r[10].strip(),
            "type": r[11].strip()
        })
wb_target.close()
print(f"Loaded {len(empty_rows)} empty rows.")

# 2. Build index from master DB: data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx
master_path = 'data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx'
wb_master = openpyxl.load_workbook(master_path, read_only=True, data_only=True)

master_by_name = {}
master_by_phone = {}

for sname in wb_master.sheetnames:
    ws = wb_master[sname]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        r = [str(c) if c is not None else "" for c in row]
        # in sheet "기업 정보":
        # col 2: 회사명, col 3: 사업자번호, col 4: 대표자, col 5: 우편번호, col 6: 주소, col 7: 전화번호, col 12: 설립일자
        c_name = r[2].strip() if len(r) > 2 else ""
        c_biz = r[3].strip() if len(r) > 3 else ""
        c_rep = r[4].strip() if len(r) > 4 else ""
        c_addr = r[6].strip() if len(r) > 6 else ""
        c_phone = r[7].strip() if len(r) > 7 else ""
        c_est = r[12].strip() if len(r) > 12 else ""
        
        # clean phone
        clean_p = re.sub(r'[^0-9]', '', c_phone)
        clean_n = re.sub(r'[\s\(\)주식회사㈜]', '', c_name)
        
        info = {"name": c_name, "biz": c_biz, "rep": c_rep, "addr": c_addr, "phone": c_phone, "est": c_est, "sheet": sname, "row": i+1}
        
        if clean_n and c_biz:
            if clean_n not in master_by_name:
                master_by_name[clean_n] = []
            master_by_name[clean_n].append(info)
            
        if clean_p and len(clean_p) >= 7 and c_biz:
            if clean_p not in master_by_phone:
                master_by_phone[clean_p] = []
            master_by_phone[clean_p].append(info)

wb_master.close()
print(f"Master indexed: {len(master_by_name)} names, {len(master_by_phone)} phones.")

# 3. Match
matched_by_name = 0
matched_by_phone = 0
matches = []

for item in empty_rows:
    cn = re.sub(r'[\s\(\)주식회사㈜]', '', item["name"])
    cp = re.sub(r'[^0-9]', '', item["phone"])
    m_info = None
    match_type = ""
    
    if cn and cn in master_by_name:
        m_info = master_by_name[cn][0]
        match_type = "name"
        matched_by_name += 1
    elif cp and len(cp) >= 7 and cp in master_by_phone:
        m_info = master_by_phone[cp][0]
        match_type = "phone"
        matched_by_phone += 1
        
    if m_info:
        matches.append({
            "target": item,
            "matched_by": match_type,
            "master": m_info
        })

print(f"Matches found: {len(matches)} (by name: {matched_by_name}, by phone: {matched_by_phone})")
with open("C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/master_matches.json", "w", encoding="utf-8") as f:
    json.dump(matches, f, ensure_ascii=False, indent=2)

import openpyxl
import json

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
sheet = wb[wb.sheetnames[0]]

names = ['브이알코리아', '블랙홀엔터테인먼트', '블루레코딩', '비문', '에프지엔터테인먼트', '우수경', '희웹툰', '작은통일출판사']
found = []

for i, row in enumerate(sheet.iter_rows(values_only=True)):
    r = [str(c) if c is not None else "" for c in row]
    r_joined = " ".join(r)
    for n in names:
        if n in r_joined:
            found.append({
                "row_idx": i+1,
                "순번": r[1],
                "사업체명": r[2],
                "사업자번호": r[3],
                "설립연도": r[4],
                "설립월": r[5],
                "대표자": r[6],
                "연락처": r[7],
                "시군구": r[9],
                "주소": r[10],
                "업종": r[11],
                "비고": r[14]
            })

with open("C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/found_prev.json", "w", encoding="utf-8") as f:
    json.dump(found, f, ensure_ascii=False, indent=2)

wb.close()
print("Saved found items, count:", len(found))

import openpyxl
import json
import os

files = [
    'data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx',
    'data/배현명_실사DB_16컬럼_전수검증완료_최종본.xlsx'
]

targets = ['우수경', '이석창']

results = []

for fpath in files:
    if not os.path.exists(fpath):
        print(f"File not found: {fpath}")
        continue
    print(f"Scanning {fpath}...")
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        for sname in wb.sheetnames:
            print(f"  Sheet: {sname}")
            ws = wb[sname]
            headers = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c) if c is not None else "" for c in row]
                    continue
                row_str = [str(c) if c is not None else "" for c in row]
                row_joined = " | ".join(row_str)
                for t in targets:
                    if t in row_joined:
                        res = {"file": fpath, "sheet": sname, "row_idx": i+1, "matched": t, "row": row_str, "headers": headers}
                        results.append(res)
                        print(f"    FOUND {t} in row {i+1}: {row_str}")
        wb.close()
    except Exception as e:
        print(f"Error reading {fpath}: {e}")

print("Total matches:", len(results))
with open("C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/target_matches.json", "w", encoding="utf-8") as out:
    json.dump(results, out, ensure_ascii=False, indent=2)

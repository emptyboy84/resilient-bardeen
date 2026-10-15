import openpyxl
import json

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True, data_only=True)
ws = wb.sheetnames[0]
sheet = wb[ws]

headers = []
target_rows = []

for i, row in enumerate(sheet.iter_rows(values_only=True)):
    if i == 0:
        headers = [str(c) if c is not None else "" for c in row]
        continue
    # look at rows around 1085-1100
    if 1080 <= i+1 <= 1100:
        target_rows.append({"row_idx": i+1, "row": [str(c) if c is not None else "" for c in row]})

with open("C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/rows_inspect.json", "w", encoding="utf-8") as f:
    json.dump({"headers": headers, "rows": target_rows}, f, ensure_ascii=False, indent=2)

wb.close()
print("Saved inspection to rows_inspect.json")

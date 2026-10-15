import openpyxl

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path)
ws = wb[wb.sheetnames[0]]

cleaned_rep_count = 0
for row_idx in range(1091, 2315):
    cell = ws.cell(row=row_idx, column=7)
    val = str(cell.value or "").strip()
    if any(k in val for k in ["이름으로", "로그인", "조회", "더보기", "회원전용"]):
        cell.value = ""
        cleaned_rep_count += 1

wb.save(target_path)
wb.close()
print(f"Cleaned {cleaned_rep_count} representative artifacts from Excel.")

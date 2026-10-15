import openpyxl
import json
import os
import csv

targets = ['희웹툰', '010-8074-7622', '8074-7622', '80747622']

# 1. Check CSV report
csv_path = 'data/배현명_실사DB_네이버지도_카카오맵_전화번호검증리포트.csv'
if os.path.exists(csv_path):
    print("Scanning CSV:", csv_path)
    with open(csv_path, mode='r', encoding='utf-8-sig', errors='ignore') as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            row_str = " ".join(row)
            for t in targets:
                if t in row_str:
                    print(f"  CSV Match at row {i+1}: {row}")

# 2. Check Excel files
for fpath in ['data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx', 'data/배현명_실사DB_16컬럼_전수검증완료_최종본.xlsx']:
    if not os.path.exists(fpath):
        continue
    print("Scanning Excel:", fpath)
    wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            row_str = [str(c) if c is not None else "" for c in row]
            row_joined = " ".join(row_str)
            for t in targets:
                if t in row_joined:
                    print(f"  Excel Match in {sname} row {i+1} for {t}: {row_str[:15]}")
    wb.close()

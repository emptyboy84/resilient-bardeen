import openpyxl
import json
import time

t0 = time.time()
print("Opening workbook...", flush=True)
wb = openpyxl.load_workbook('data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx', read_only=True, data_only=True)
print(f"Workbook opened in {time.time()-t0:.2f}s", flush=True)

targets = ['브이알코리아', '블랙홀엔터테인먼트', '블루레코딩', '비문', '9514-1435', '349-8000', '602-81-61888', '6028161888']

for sheetname in wb.sheetnames:
    print(f"Scanning sheet: {sheetname}", flush=True)
    ws = wb[sheetname]
    matches = []
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c) if c is not None else "" for c in row]
            print("Headers:", headers, flush=True)
            continue
        row_str = [str(c) if c is not None else "" for c in row]
        row_joined = " | ".join(row_str)
        for t in targets:
            if t in row_joined:
                matches.append({"line": i+1, "matched": t, "row": row_str})
                print(f"Match for {t} at row {i+1}: {row_str}", flush=True)
                break
    with open(f"C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/matches_{sheetname}.json", "w", encoding="utf-8") as f:
        json.dump({"headers": headers, "matches": matches}, f, ensure_ascii=False, indent=2)

print("Done!", flush=True)

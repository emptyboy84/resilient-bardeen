import openpyxl

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True)
ws = wb[wb.sheetnames[0]]

total_rows_inspected = 0
has_biz_count = 0
no_biz_with_note = 0
unhandled = 0

sample_rows = [1091, 1113, 1136, 1149, 1179, 1226, 1483, 1517, 1643, 1811, 2022, 2103, 2229, 2257, 2314]
samples_output = []

for i, row in enumerate(ws.iter_rows(values_only=True)):
    row_idx = i + 1
    if 1091 <= row_idx <= 2314:
        total_rows_inspected += 1
        r = [str(c) if c is not None else "" for c in row]
        bno = r[3].strip()
        note = r[14].strip()
        
        if bno:
            has_biz_count += 1
        elif note:
            no_biz_with_note += 1
        else:
            unhandled += 1
            
        if row_idx in sample_rows:
            samples_output.append({
                "row": row_idx,
                "순번": r[1],
                "name": r[2],
                "bno": r[3],
                "est_year": r[4],
                "est_month": r[5],
                "rep": r[6],
                "close_date": r[13],
                "note": r[14]
            })

wb.close()

with open('C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/verification_report.json', 'w', encoding='utf-8') as f:
    import json
    json.dump({
        "total_rows_inspected": total_rows_inspected,
        "has_biz_count": has_biz_count,
        "no_biz_with_note": no_biz_with_note,
        "unhandled": unhandled,
        "samples": samples_output
    }, f, ensure_ascii=False, indent=2)

print(f"Inspection Summary (Rows 1091 ~ 2314, Total: {total_rows_inspected}):")
print(f"  - Filled 사업자번호: {has_biz_count}")
print(f"  - Blank 사업자번호 with 확인불가/비고: {no_biz_with_note}")
print(f"  - Completely unhandled (No biz AND No note): {unhandled}")

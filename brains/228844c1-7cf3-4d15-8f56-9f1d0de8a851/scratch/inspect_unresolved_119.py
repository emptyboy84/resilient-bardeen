import json
import openpyxl

with open('C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/resolved_by_row.json', encoding='utf-8') as f:
    resolved = json.load(f)

resolved_rows = {int(k) for k in resolved.keys()}

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"
wb = openpyxl.load_workbook(target_path, read_only=True)
ws = wb[wb.sheetnames[0]]

unresolved = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    row_idx = i + 1
    if 1091 <= row_idx <= 2314:
        r = [str(c) if c is not None else "" for c in row]
        if not r[3].strip() and row_idx not in resolved_rows:
            unresolved.append({
                "row_idx": row_idx,
                "순번": r[1],
                "name": r[2].strip(),
                "rep": r[6].strip(),
                "phone": r[7].strip(),
                "addr": r[10].strip(),
                "type": r[11].strip()
            })
wb.close()

print(f"Total remaining unresolved real rows: {len(unresolved)}")

with open('C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/unresolved_119.json', 'w', encoding='utf-8') as f:
    json.dump(unresolved, f, ensure_ascii=False, indent=2)

with open('C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/unresolved_119_list.txt', 'w', encoding='utf-8') as f:
    for idx, u in enumerate(unresolved):
        f.write(f"[{idx+1}] Row {u['row_idx']} (순번 {u['순번']}): {u['name']} | Type: {u['type']} | Rep: {u['rep']} | Phone: {u['phone']} | Addr: {u['addr']}\n")

print("Saved to scratch/unresolved_119_list.txt")

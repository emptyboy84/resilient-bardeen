import openpyxl
import json
import os

target_path = r"C:\Users\lg\Desktop\오늘\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명.xlsx"

with open('C:/Users/lg/.gemini/antigravity/brain/228844c1-7cf3-4d15-8f56-9f1d0de8a851/scratch/final_master_resolution.json', encoding='utf-8') as f:
    resolution = json.load(f)

print(f"Loaded {len(resolution)} rows to write.")

wb = openpyxl.load_workbook(target_path)
ws = wb[wb.sheetnames[0]]

updated_biz_count = 0
updated_note_count = 0

for row_str, data in resolution.items():
    row_idx = int(row_str)
    
    # Col 4: 사업자번호
    biz = data.get("biz", "").strip()
    if biz:
        ws.cell(row=row_idx, column=4, value=biz)
        updated_biz_count += 1
        
    # Col 5: 설립(년)
    est_y = data.get("est_year", "").strip()
    curr_y = ws.cell(row=row_idx, column=5).value
    if est_y and (curr_y is None or not str(curr_y).strip()):
        ws.cell(row=row_idx, column=5, value=est_y)
        
    # Col 6: 설립(월)
    est_m = data.get("est_month", "").strip()
    curr_m = ws.cell(row=row_idx, column=6).value
    if est_m and (curr_m is None or not str(curr_m).strip()):
        ws.cell(row=row_idx, column=6, value=est_m)
        
    # Col 7: 대표자 성함
    rep = data.get("rep", "").strip()
    curr_rep = ws.cell(row=row_idx, column=7).value
    if rep and (curr_rep is None or not str(curr_rep).strip()):
        ws.cell(row=row_idx, column=7, value=rep)
        
    # Col 14: 폐업 날짜
    close_d = data.get("close_date", "").strip()
    curr_close = ws.cell(row=row_idx, column=14).value
    if close_d and (curr_close is None or not str(curr_close).strip()):
        ws.cell(row=row_idx, column=14, value=close_d)
        
    # Col 15: 비고
    note = data.get("note", "").strip()
    curr_note = ws.cell(row=row_idx, column=15).value
    if note:
        if curr_note is None or not str(curr_note).strip():
            ws.cell(row=row_idx, column=15, value=note)
            updated_note_count += 1
        else:
            # If already has note and note is different, append or keep
            curr_str = str(curr_note).strip()
            if note not in curr_str:
                ws.cell(row=row_idx, column=15, value=f"{curr_str} / {note}")
                updated_note_count += 1

wb.save(target_path)
wb.close()

print(f"Successfully saved to {target_path}!")
print(f"  - Updated 사업자번호: {updated_biz_count} rows")
print(f"  - Updated 비고/확인불가: {updated_note_count} rows")

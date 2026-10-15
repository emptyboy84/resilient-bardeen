import shutil
import os
import datetime
import openpyxl
import json

file_path = r"C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx"
backup_dir = r"C:\Users\lg\Desktop\Today"
timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = os.path.join(backup_dir, f"(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명__BACKUP_{timestamp}.xlsx")

# 1. Create Backup
shutil.copy2(file_path, backup_path)
print(f"Backup created successfully: {backup_path}")

# 2. Extract 318 target rows
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["기업체 리스트"]

targets = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if not any(v is not None and str(v).strip() != "" for v in vals):
        continue
    biz_no = vals[3]
    if not biz_no or str(biz_no).strip() == "":
        targets.append({
            "row": r,
            "gubun": vals[0],
            "seq": vals[1],
            "name": str(vals[2]).strip() if vals[2] is not None else "",
            "biz_no": "",
            "est_year": vals[4],
            "est_month": vals[5],
            "ceo": str(vals[6]).strip() if vals[6] is not None else "",
            "tel": str(vals[7]).strip() if vals[7] is not None else "",
            "fax": str(vals[8]).strip() if vals[8] is not None else "",
            "sigungu": str(vals[9]).strip() if vals[9] is not None else "",
            "addr": str(vals[10]).strip() if vals[10] is not None else "",
            "industry": str(vals[11]).strip() if vals[11] is not None else "",
            "contact_res": vals[12],
            "closed_date": vals[13],
            "remark": str(vals[14]).strip() if vals[14] is not None else "",
            "emp_cnt": vals[15],
            "sales_24": vals[16],
            "sales_25": vals[17]
        })

print(f"Extracted {len(targets)} target companies.")
json_path = r"C:\Users\lg\.gemini\antigravity\brain\cc5f8116-f40f-4d66-a990-d85efe7bc64f\scratch\target_318_companies.json"
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(targets, f, ensure_ascii=False, indent=2)

print(f"Saved target companies to {json_path}")

import openpyxl
import sys

# Ensure stdout handles utf-8 if possible
sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx"

wb = openpyxl.load_workbook(file_path, data_only=True)

with open(r"C:\Users\lg\.gemini\antigravity\brain\cc5f8116-f40f-4d66-a990-d85efe7bc64f\scratch\analysis_summary.txt", "w", encoding="utf-8") as out:
    for name in wb.sheetnames:
        ws = wb[name]
        out.write(f"\n==================================================\n")
        out.write(f"Sheet: {name} (max_row={ws.max_row}, max_column={ws.max_column})\n")
        out.write(f"==================================================\n")
        headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        out.write(f"Headers: {headers}\n\n")
        
        col_missing = {h: 0 for h in headers if h}
        total_rows = ws.max_row - 1
        
        for row in range(2, ws.max_row + 1):
            for col_idx, h in enumerate(headers, start=1):
                if h:
                    val = ws.cell(row, col_idx).value
                    if val is None or str(val).strip() == "":
                        col_missing[h] += 1
                        
        out.write(f"Total data rows: {total_rows}\n")
        out.write("Missing counts:\n")
        for h, cnt in col_missing.items():
            out.write(f"  [{h}]: {cnt} / {total_rows} missing ({cnt/total_rows*100:.1f}%)\n")

print("Analysis written to analysis_summary.txt")

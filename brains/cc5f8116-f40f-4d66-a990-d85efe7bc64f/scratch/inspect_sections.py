import openpyxl

file_path = r"C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["기업체 리스트"]

# Check rows where 구분 changes or blank
sections = {}
for r in range(2, ws.max_row + 1):
    gubun = ws.cell(r, 1).value
    biz_no = ws.cell(r, 4).value
    name = ws.cell(r, 3).value
    if name is not None or biz_no is not None or gubun is not None:
        g_str = str(gubun)
        if g_str not in sections:
            sections[g_str] = {"total": 0, "missing_biz_no": 0, "rows": []}
        sections[g_str]["total"] += 1
        if not biz_no or str(biz_no).strip() == "":
            sections[g_str]["missing_biz_no"] += 1
            sections[g_str]["rows"].append(r)

with open(r"C:\Users\lg\.gemini\antigravity\brain\cc5f8116-f40f-4d66-a990-d85efe7bc64f\scratch\sections_analysis.txt", "w", encoding="utf-8") as f:
    for g, data in sections.items():
        f.write(f"구분: [{g}] -> Total: {data['total']}, Missing 사업자번호: {data['missing_biz_no']}\n")
        if data["missing_biz_no"] > 0:
            f.write(f"  First 5 row indices: {data['rows'][:5]}\n")
            f.write(f"  Last 5 row indices: {data['rows'][-5:]}\n")

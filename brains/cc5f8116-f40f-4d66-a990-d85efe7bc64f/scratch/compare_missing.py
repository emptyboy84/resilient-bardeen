import openpyxl

file_path = r"C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["기업체 리스트"]

headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]

missing_biz_rows = []
present_biz_rows = []

for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if not any(v is not None and str(v).strip() != "" for v in vals):
        continue
    biz_no = vals[3]
    if not biz_no or str(biz_no).strip() == "":
        missing_biz_rows.append(vals)
    else:
        present_biz_rows.append(vals)

with open(r"C:\Users\lg\.gemini\antigravity\brain\cc5f8116-f40f-4d66-a990-d85efe7bc64f\scratch\compare_missing.txt", "w", encoding="utf-8") as f:
    f.write(f"Total rows with 사업자번호 missing: {len(missing_biz_rows)}\n")
    f.write(f"Total rows with 사업자번호 present: {len(present_biz_rows)}\n\n")
    
    f.write("--- Missing counts in rows where 사업자번호 is MISSING (N=318) ---\n")
    for col_idx, h in enumerate(headers):
        if not h: continue
        missing_cnt = sum(1 for row in missing_biz_rows if row[col_idx] is None or str(row[col_idx]).strip() == "")
        f.write(f"  [{h}]: {missing_cnt} / {len(missing_biz_rows)} ({missing_cnt/len(missing_biz_rows)*100:.1f}% missing)\n")
        
    f.write("\n--- Missing counts in rows where 사업자번호 is PRESENT (N=1995) ---\n")
    for col_idx, h in enumerate(headers):
        if not h: continue
        missing_cnt = sum(1 for row in present_biz_rows if row[col_idx] is None or str(row[col_idx]).strip() == "")
        f.write(f"  [{h}]: {missing_cnt} / {len(present_biz_rows)} ({missing_cnt/len(present_biz_rows)*100:.1f}% missing)\n")

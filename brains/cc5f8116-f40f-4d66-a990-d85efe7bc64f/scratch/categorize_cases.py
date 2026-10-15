import openpyxl

file_path = r"C:\Users\lg\Desktop\Today\(안상숙)(STI)_(모집단현행화)_부산정보산업진흥원_2026_부산_CT산업_실태조사_260910_크레탑_배현명_.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["기업체 리스트"]

missing_biz_rows = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if not any(v is not None and str(v).strip() != "" for v in vals):
        continue
    biz_no = vals[3]
    if not biz_no or str(biz_no).strip() == "":
        missing_biz_rows.append({
            "row": r,
            "seq": vals[1],
            "name": vals[2],
            "ceo": vals[6],
            "tel": vals[7],
            "sigungu": vals[9],
            "addr": vals[10],
            "industry": vals[11],
            "remark": vals[14]
        })

with open(r"C:\Users\lg\.gemini\antigravity\brain\cc5f8116-f40f-4d66-a990-d85efe7bc64f\scratch\sample_cases.txt", "w", encoding="utf-8") as f:
    f.write(f"Total target companies: {len(missing_biz_rows)}\n\n")
    # Category 1: Has Name + Tel + Addr (Easiest to match)
    cat1 = [c for c in missing_biz_rows if c['tel'] and c['addr']]
    f.write(f"1. Has Tel & Addr: {len(cat1)} cases\n")
    for c in cat1[:5]:
        f.write(f"   Row {c['row']} | {c['name']} | {c['tel']} | {c['addr']}\n")
        
    # Category 2: Has Name + Addr only (no tel)
    cat2 = [c for c in missing_biz_rows if not c['tel'] and c['addr']]
    f.write(f"\n2. Has Addr only (No Tel): {len(cat2)} cases\n")
    for c in cat2[:5]:
        f.write(f"   Row {c['row']} | {c['name']} | {c['addr']}\n")
        
    # Category 3: Has Name + Tel only (no addr)
    cat3 = [c for c in missing_biz_rows if c['tel'] and not c['addr']]
    f.write(f"\n3. Has Tel only (No Addr): {len(cat3)} cases\n")
    for c in cat3[:5]:
        f.write(f"   Row {c['row']} | {c['name']} | {c['tel']} | CEO: {c['ceo']}\n")
        
    # Category 4: Name only
    cat4 = [c for c in missing_biz_rows if not c['tel'] and not c['addr']]
    f.write(f"\n4. Name only (No Tel & No Addr): {len(cat4)} cases\n")
    for c in cat4[:5]:
        f.write(f"   Row {c['row']} | {c['name']} | CEO: {c['ceo']}\n")

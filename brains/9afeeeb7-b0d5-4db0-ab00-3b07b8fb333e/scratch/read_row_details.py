import openpyxl

excel_path = r"data/(데이터검증팀)실사 DB 취합_전달용_260720_배현명.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb.active

headers = [cell.value for cell in ws[1]]
row_vals = [cell.value for cell in ws[20418]]

for i, (h, v) in enumerate(zip(headers, row_vals), start=1):
    print(f"Col {i:02d} | {h}: {v}")

wb.close()
